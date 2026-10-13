"""Read every PDF of a folder with one Reader, export each document, and write one workbook with a row for each.

This is the script of docs/how-to/batch-a-folder.md. Run it from the repository folder:

    .venv/bin/python examples/batch_folder.py [FOLDER [DESTINATION]]

With no arguments it reads the three samples of tests/fixtures and writes into out/: one subfolder for each document and
out/documents.xlsx with one row for each.
"""

import sys
from pathlib import Path

import openpyxl

import extract

HEADER = [
    "File",
    "Status",
    "Unreadable pages",
    "Words to review",
    "Invoice number",
    "Total",
    "Total (value)",
    "Total status",
]


def value_of(doc, cell):
    """The exact value of an invoice field as a Decimal: the cell holds the text, its number holds the value."""
    return next((number.value for number in doc.numbers if number.word_ids == cell.word_ids), None)


def row_of(doc):
    invoice = doc.invoice
    total = invoice.total if invoice else None
    return [
        doc.name,
        doc.status,
        ", ".join(str(page) for page in doc.unreadable_pages) or None,
        len(doc.doubtful()),
        invoice.number.text if invoice and invoice.number else None,
        total.text if total else None,
        value_of(doc, total) if total else None,
        doc.field_status(total) if total else "not found",
    ]


def read_folder(folder, destination):
    """Read the PDFs of `folder` in name order, export each into its own subfolder, and write one workbook."""
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Documents"
    sheet.append(HEADER)
    with extract.Reader(extract.Options(work_dir="work")) as reader:
        for pdf in sorted(Path(folder).glob("*.pdf")):
            try:
                doc = reader.read(pdf)
            except (extract.UnsupportedFile, extract.PasswordProtected) as error:
                print(f"{pdf.name}: not read ({error.code})")
                continue
            doc.export(Path(destination) / pdf.stem)
            sheet.append(row_of(doc))
    Path(destination).mkdir(parents=True, exist_ok=True)
    book.save(Path(destination) / "documents.xlsx")


if __name__ == "__main__":
    folder = sys.argv[1] if len(sys.argv) > 1 else "tests/fixtures"
    destination = sys.argv[2] if len(sys.argv) > 2 else "out"
    read_folder(folder, destination)
    for row in openpyxl.load_workbook(Path(destination) / "documents.xlsx")["Documents"].iter_rows(values_only=True):
        print(row)
