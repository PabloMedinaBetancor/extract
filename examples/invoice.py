"""Follow an invoice whose total is 30 cents off, from the failed number to the fixed workbook.

This is the script of docs/tutorials/invoice.md. Run it from the repository folder:

    .venv/bin/python examples/invoice.py

It reads tests/fixtures/mismatch.pdf, decides the three doubtful words and the total, and writes the workbook into out/.
The decisions are logged in work/corrections.jsonl, and the cut-out of the total is saved as total.png in the current
folder.
"""

import json
from collections import Counter
from pathlib import Path

import openpyxl

import extract


def main():
    options = extract.Options(work_dir="work")
    doc = extract.read("tests/fixtures/mismatch.pdf", options=options)
    print(doc.name, doc.status, doc.doc_type, doc.locale)

    for word in doc.doubtful():
        reasons = [reason.key.rsplit(".", 1)[-1] for reason in word.reasons]
        print(word.id, word.text, round(word.p_wrong, 2), word.suggestion, reasons)
    for word in doc.doubtful():
        doc.accept(word.id)
    print(len(doc.doubtful()), doc.word("p1w11").status, doc.word("p1w11").review)

    print(Counter(number.status for number in doc.numbers))
    for number in doc.numbers:
        if number.status == "mismatch":
            print(number.text, [(rule.key.rsplit(".", 1)[-1], rule.result) for rule in number.rules])

    invoice = doc.invoice
    for line in invoice.lines:
        print(line.description.text, line.amount.text, doc.field_status(line.amount))
    print(invoice.subtotal.text, doc.field_status(invoice.subtotal))
    print(invoice.total.text, doc.field_status(invoice.total))

    total = next(number for number in doc.numbers if number.text == "101.913,81")
    failed = next(rule for rule in total.rules if rule.result == "fail")
    print(failed.key, failed.params["expected"], failed.params["found"])
    print(total.suggestion)

    Path("total.png").write_bytes(doc.word_crop("p1w75"))

    doc.use_suggestion("p1w75")
    word = doc.word("p1w75")
    print(word.text, word.original, word.review, word.status)
    print(doc.field_status(doc.invoice.total), doc.field_status(doc.invoice.subtotal))
    print(Counter(number.status for number in doc.numbers))

    last = json.loads(Path("work/corrections.jsonl").read_text(encoding="utf-8").splitlines()[-1])
    print({key: last[key] for key in ("word", "original", "corrected", "action")})

    path = doc.export("out/", ("xlsx",))[0]
    book = openpyxl.load_workbook(path)
    print(book.sheetnames)
    for row in book["Fields"].iter_rows(values_only=True):
        if row[0] == "total":
            print(row[:4])
    for row in book["Corrections"].iter_rows(min_row=2, values_only=True):
        print(row[:5])


if __name__ == "__main__":
    main()
