import ast
import dataclasses
import inspect as inspect_module
import json
import os
import sys
import tempfile
import time
import unittest
from pathlib import Path
from unittest import mock

import openpyxl
import pikepdf

import extract
from extract import events
from extract.api import Reader
from tests.fake_fix_guard import FixGuard, fix
from tests.fake_guard import DoubtingGuard, FailingGuard, FakeGuard, PartialGuard
from tests.support import (
    SANS,
    SPANISH_LINES,
    SetLexicon,
    broken_page_pdf,
    digital_pdf,
    mixed_pdf,
    scanned_pdf,
    text_page_image,
)

FAKE = "tests.fake_engine:load_engine"
LEXICON = SetLexicon({w.lower().strip(".,:") for line in SPANISH_LINES for w in line.split() if len(w) > 2})
PINNED_ALL = [
    "CancelToken", "Cancelled", "Cell", "DigitSuggestion", "Document", "DocumentRead", "EngineLoadError", "Event",
    "Exported", "ExtractError", "FileInfo", "Finished", "Fix", "Invoice", "InvoiceLine", "ModelMismatch", "Number", "Options",
    "OutputExists", "Page", "PageRead", "PasswordProtected", "Progress", "Reader", "Reading", "Reason", "Rule", "SourceChanged",
    "Started", "UnsupportedFile", "VatPair", "Warning", "Word", "inspect", "load", "read",
]  # fmt: skip
PINNED_SIGNATURES = {
    "read": "(path: 'str | Path', options: 'Options | None' = None, on_event: 'Callable[[events.Event], None] | None' = None, cancel: 'CancelToken | None' = None, *, correct: 'str | None' = None) -> 'Document'",
    "load": "(path: 'str | Path', source: 'str | Path | None' = None, options: 'Options | None' = None) -> 'Document'",
    "inspect": "(path: 'str | Path') -> 'FileInfo'",
    "Reader.__init__": "(self, options: 'Options | None' = None, *, guard: 'Any' = None, lexicon: 'Any' = None) -> 'None'",
    "Reader.read": "(self, path: 'str | Path', on_event: 'Callable[[events.Event], None] | None' = None, cancel: 'CancelToken | None' = None, *, correct: 'str | None' = None) -> 'Document'",
    "Document.set_mode": "(self, mode: 'str') -> 'None'",
    "Document.undo": "(self, word_id: 'str') -> 'None'",
    "Document.undo_page": "(self, number: 'int') -> 'int'",
    "Document.undo_all": "(self) -> 'int'",
    "Document.tokens": "(self, page: 'int') -> 'tuple[Word, ...]'",
    "Reader.close": "(self) -> 'None'",
    "Document.correct": "(self, word_id: 'str', text: 'str') -> 'None'",
    "Document.accept": "(self, word_id: 'str') -> 'None'",
    "Document.use_suggestion": "(self, word_id: 'str') -> 'None'",
    "Document.doubtful": "(self) -> 'list[Word]'",
    "Document.save": "(self, path: 'str | Path') -> 'None'",
    "Document.export": "(self, destination: 'str | Path', formats: 'Sequence[str]' = ('pdf', 'xlsx', 'json', 'txt'), *, name: 'str | None' = None, language: 'str' = 'en', on_event: 'Callable[[events.Event], None] | None' = None, cancel: 'CancelToken | None' = None) -> 'list[Path]'",
    "Document.page_image": "(self, number: 'int', zoom: 'float' = 1.0) -> 'bytes'",
    "Document.word_crop": "(self, word_id: 'str', pad: 'int' = 8) -> 'bytes'",
}  # fmt: skip


def signature_of(name):
    target = extract
    for part in name.split("."):
        target = getattr(target, part)
    return str(inspect_module.signature(target))


class Flag:
    """The cancel token of the tests, set from an event handler."""

    def __init__(self):
        self.value = False

    def is_set(self):
        return self.value


class ApiTest(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        patcher = mock.patch.object(Reader, "_engine_factory", FAKE)
        patcher.start()
        self.addCleanup(patcher.stop)
        self.options = extract.Options(work_dir=self.tmp / "work", processes=1, threads=1)
        self.digital = digital_pdf(self.tmp / "digital.pdf", [SPANISH_LINES, SPANISH_LINES[:3]])
        scan = text_page_image(SPANISH_LINES, SANS, 24, 1240, 1754)
        self.scanned = scanned_pdf(self.tmp / "scanned.pdf", [scan, scan])
        self.mixed = mixed_pdf(self.tmp / "mixed.pdf", SPANISH_LINES, scan)

    def reader(self, guard=None, **kwargs):
        reader = Reader(self.options, guard=guard if guard is not None else FakeGuard(), lexicon=LEXICON, **kwargs)
        self.addCleanup(reader.close)
        return reader

    def read(self, path, guard=None, **kwargs):
        return self.reader(guard).read(path, **kwargs)


class SurfaceTest(ApiTest):
    def test_A1_the_public_names_are_pinned(self):
        self.assertEqual(sorted(extract.__all__), PINNED_ALL)
        for name in PINNED_ALL:
            self.assertTrue(hasattr(extract, name), name)

    def test_A1_the_signatures_are_pinned(self):
        for name, expected in PINNED_SIGNATURES.items():
            with self.subTest(name=name):
                self.assertEqual(signature_of(name), expected)

    def test_A1_options_fields_are_pinned(self):
        fields = [(f.name, f.default) for f in dataclasses.fields(extract.Options)]
        self.assertEqual(
            fields,
            [
                ("languages", ("es", "en")),
                ("review", "balanced"),
                ("recycle_pages", 120),
                ("work_dir", None),
                ("processes", None),
                ("threads", None),
                ("backend", "auto"),
                ("correct", "off"),
                ("second_reading", False),
            ],
        )

    def test_A1_the_guard_and_the_rules_are_not_public(self):
        for name in ("Guard", "checks", "Scored", "DocumentData"):
            self.assertNotIn(name, extract.__all__)


class ReadTest(ApiTest):
    def test_A2_read_and_reader_read_give_equal_documents(self):
        for path in (self.digital, self.scanned, self.mixed):
            with self.subTest(path=path.name):
                with (
                    mock.patch.object(extract.api, "_default_guard", lambda: FakeGuard()),
                    mock.patch.object(extract.api, "_default_lexicon", lambda languages: LEXICON),
                ):
                    alone = extract.read(path, self.options)
                self.assertEqual(alone, self.read(path))

    def test_A2_three_documents_are_exported_in_four_formats_and_the_files_reopen(self):
        import pypdfium2 as pdfium

        reader = self.reader()
        for path in (self.digital, self.scanned, self.mixed):
            with self.subTest(path=path.name):
                document = reader.read(path)
                out = self.tmp / "out" / path.stem
                files = {p.name.split(".", 1)[1]: p for p in document.export(out)}
                self.assertEqual(sorted(files), ["extract.json", "searchable.pdf", "txt", "xlsx"])
                with pikepdf.open(files["searchable.pdf"]) as pdf:
                    self.assertEqual(len(pdf.pages), len(document.pages))
                self.assertEqual(extract.load(files["extract.json"], path), document)
                self.assertIn("Numbers", openpyxl.load_workbook(files["xlsx"]).sheetnames)
                self.assertEqual(len(files["txt"].read_text(encoding="utf-8").split("\f")), len(document.pages))
                pdfium.PdfDocument(str(files["searchable.pdf"])).close()

    def test_A2_a_mixed_document_has_native_and_read_pages(self):
        document = self.read(self.mixed)
        self.assertEqual([p.status for p in document.pages], ["native", "ocr"])
        self.assertEqual({w.status for w in document.pages[0].words}, {"native"})
        self.assertEqual({w.status for w in document.pages[1].words}, {"ok"})

    def test_A2_without_a_guard_the_words_are_unscored_and_the_user_is_told(self):
        seen = []
        reader = Reader(self.options, guard=False, lexicon=LEXICON)
        self.addCleanup(reader.close)
        document = reader.read(self.scanned, on_event=seen.append)
        self.assertEqual({w.status for w in document.words()}, {"unscored"})
        warnings = [e for e in seen if isinstance(e, events.Warning)]
        self.assertEqual({e.params["cause"] for e in warnings}, {"no_guard"})
        self.assertEqual(len(document.doubtful()), sum(len(p.words) for p in document.pages))

    def test_A2_words_the_guard_did_not_judge_are_unscored_and_one_warning_counts_them(self):
        seen = []
        document = self.read(self.scanned, PartialGuard(10), on_event=seen.append)
        first = document.pages[0].words
        self.assertEqual({w.status for w in first[:10]}, {"ok"})
        self.assertEqual({w.status for w in first[10:]}, {"unscored"})
        warnings = [e for e in seen if isinstance(e, events.Warning) and e.params["page"] == 1]
        self.assertEqual(len(warnings), 1)
        self.assertEqual(warnings[0].params["words"], len(first) - 10)

    def test_A2_the_review_levels_reach_the_guard_as_its_budgets(self):
        for review, budget in (("fewer", 0.03), ("balanced", 0.05), ("more", 0.08)):
            with self.subTest(review=review):
                guard = FakeGuard()
                options = extract.Options(work_dir=self.tmp / f"w-{review}", processes=1, threads=1, review=review)
                reader = Reader(options, guard=guard, lexicon=LEXICON)
                self.addCleanup(reader.close)
                reader.read(self.scanned)
                self.assertEqual({c["budget"] for c in guard.calls}, {budget})
                self.assertEqual({c["locale"] for c in guard.calls}, {"auto"})
                self.assertEqual({c["languages"] for c in guard.calls}, {("es", "en")})

    def test_A3_a_suggestion_that_no_longer_fits_the_text_is_refused_not_applied(self):
        document = self.read(self.scanned, DoubtingGuard({"word3": None}))
        word = next(w for w in document.doubtful() if w.text == "word3")
        self.assertIsNone(word.suggestion)
        document._data = dataclasses.replace(
            document._data,
            numbers=(
                extract.Number(
                    (word.id,), "word3", None, "amount", "mismatch", (), extract.DigitSuggestion(4, "9", "1")
                ),
            ),
        )
        with self.assertRaises(ValueError):
            document.use_suggestion(word.id)
        self.assertEqual(document.word(word.id).text, "word3")

    def test_A2_a_guard_that_fails_leaves_unscored_words_and_a_warning_not_an_exception(self):
        seen = []
        document = self.read(self.scanned, FailingGuard(), on_event=seen.append)
        self.assertEqual({w.status for w in document.words()}, {"unscored"})
        warning = next(e for e in seen if isinstance(e, events.Warning))
        self.assertEqual((warning.code, warning.params["cause"]), ("guard.warning.unscored", "error"))

    def test_A2_an_unreadable_page_is_partial_with_unscored_words_carrying_the_reason(self):
        with mock.patch.dict(os.environ, {"EXTRACT_FAKE_PAGE": str(self._few_words_page())}):
            document = self.read(self.scanned)
        self.assertEqual(document.status, "partial")
        self.assertEqual(document.unreadable_pages, (1, 2))
        word = document.pages[0].words[0]
        self.assertEqual(word.status, "unscored")
        self.assertEqual(word.reasons[0].key, "page.reason.few_words")

    def _few_words_page(self):
        path = self.tmp / "few.json"
        path.write_text(json.dumps([[["alpha", 0.9, [0, 0, 10, 10]], ["beta", 0.9, [20, 0, 30, 10]]]]))
        return path

    def test_A2_a_broken_page_makes_the_document_partial_and_the_rest_is_read(self):
        document = self.read(broken_page_pdf(self.tmp / "broken.pdf"))
        self.assertEqual((document.status, document.unreadable_pages), ("partial", (2,)))
        self.assertEqual([p.status for p in document.pages], ["native", "unreadable", "native"])

    def test_A2_the_packaged_word_lists_decide_which_text_layers_are_legible(self):
        reader = Reader(self.options, guard=FakeGuard())
        self.addCleanup(reader.close)
        self.assertEqual([p.status for p in reader.read(self.digital).pages], ["native", "native"])
        garbage = digital_pdf(self.tmp / "garbage.pdf", [["xqzv wkjp zzqx " * 8] * 6])
        self.assertEqual(reader.read(garbage).pages[0].status, "ocr")

    def test_A2_inspect_tells_digital_scanned_and_mixed_apart_without_reading(self):
        self.assertEqual(
            [extract.inspect(p).kind for p in (self.digital, self.scanned, self.mixed)], ["digital", "scanned", "mixed"]
        )
        info = extract.inspect(self.mixed)
        self.assertEqual((info.pages, info.digital_pages, info.scanned_pages), (2, 1, 1))
        self.assertEqual(info.size, self.mixed.stat().st_size)


class CorrectionTest(ApiTest):
    def setUp(self):
        super().setUp()
        self.guard = DoubtingGuard({"word3": "word-three", "word10": None, "word20": "twenty"})
        self.document = self.read(self.scanned, self.guard)

    def test_A3_doubtful_words_come_likeliest_first_and_corrections_leave_the_list(self):
        words = self.document.doubtful()
        self.assertEqual(sorted(w.text for w in words if w.page == 1), ["word10", "word20", "word3"])
        first = words[0].id
        self.document.accept(first)
        self.assertNotIn(first, [w.id for w in self.document.doubtful()])
        self.assertEqual(self.document.word(first).review, "accepted")

    def test_A3_a_correction_keeps_the_original_and_marks_the_user(self):
        word = next(w for w in self.document.doubtful() if w.text == "word3")
        self.document.correct(word.id, "Tres")
        changed = self.document.word(word.id)
        self.assertEqual(
            (changed.text, changed.original, changed.review, changed.corrected_by),
            ("Tres", "word3", "corrected", "user"),
        )
        self.assertIsNone(changed.suggestion)
        self.document.correct(word.id, "Tres!")
        self.assertEqual(self.document.word(word.id).original, "word3")

    def test_A3_use_suggestion_applies_it_and_a_word_without_one_is_refused(self):
        with_suggestion = next(w for w in self.document.doubtful() if w.text == "word3")
        self.document.use_suggestion(with_suggestion.id)
        self.assertEqual(self.document.word(with_suggestion.id).text, "word-three")
        without = next(w for w in self.document.doubtful() if w.text == "word10")
        with self.assertRaises(ValueError):
            self.document.use_suggestion(without.id)

    def test_A3_save_and_load_give_the_same_document_and_the_same_export(self):
        word = self.document.doubtful()[0]
        self.document.correct(word.id, "fixed")
        self.document.accept(self.document.doubtful()[0].id)
        path = self.tmp / "doc.extract.json"
        self.document.save(path)
        loaded = extract.load(path, self.scanned)
        self.assertEqual(loaded, self.document)
        before = {p.name: p for p in self.document.export(self.tmp / "a")}
        after = {p.name: p for p in loaded.export(self.tmp / "b")}
        self.assertEqual(before.keys(), after.keys())
        self.assertEqual(
            before["scanned.txt"].read_text(encoding="utf-8"), after["scanned.txt"].read_text(encoding="utf-8")
        )
        self.assertEqual(
            before["scanned.extract.json"].read_text(encoding="utf-8"),
            after["scanned.extract.json"].read_text(encoding="utf-8"),
        )

        def sheets(path):
            return [[c.value for c in row] for ws in openpyxl.load_workbook(path).worksheets for row in ws.iter_rows()]

        self.assertEqual(sheets(before["scanned.xlsx"]), sheets(after["scanned.xlsx"]))
        self.assertIn("fixed", after["scanned.txt"].read_text(encoding="utf-8"))

    def test_A3_corrections_are_logged_for_a_later_retraining_and_nothing_else_of_the_page(self):
        word = self.document.doubtful()[0]
        self.document.correct(word.id, "fixed")
        self.document.accept(self.document.doubtful()[0].id)
        lines = [json.loads(line) for line in (self.tmp / "work" / "corrections.jsonl").read_text().splitlines()]
        self.assertEqual([l["action"] for l in lines], ["correct", "accept"])  # noqa: E741
        self.assertEqual(set(lines[0]), {"sha256", "word", "original", "corrected", "action", "time"})
        self.assertEqual(lines[0]["sha256"], self.document.sha256)

    def test_A3_a_corrected_word_is_never_exported_as_doubtful(self):
        word = self.document.doubtful()[0]
        self.document.correct(word.id, "fixed")
        paths = self.document.export(self.tmp / "out", ["xlsx"])
        texts = [row[1].value for row in openpyxl.load_workbook(paths[0])["Review"].iter_rows(min_row=2)]
        self.assertNotIn("fixed", texts)


class PagesTest(ApiTest):
    def test_A2_page_images_and_crops_come_from_the_source_at_the_pages_resolution(self):
        document = self.read(self.scanned)
        png = document.page_image(1)
        self.assertTrue(png.startswith(b"\x89PNG"))
        import io

        from PIL import Image

        page = document.pages[0]
        self.assertAlmostEqual(Image.open(io.BytesIO(png)).width, page.width_px, delta=2)
        self.assertAlmostEqual(
            Image.open(io.BytesIO(document.page_image(1, zoom=0.5))).width, page.width_px / 2, delta=2
        )
        word = page.words[3]
        crop = Image.open(io.BytesIO(document.word_crop(word.id, pad=8)))
        self.assertAlmostEqual(crop.width, word.box[2] - word.box[0] + 17, delta=3)
        self.assertAlmostEqual(crop.height, word.box[3] - word.box[1] + 17, delta=3)

    def test_A5_a_source_that_changed_is_refused(self):
        document = self.read(self.scanned)
        self.scanned.write_bytes(self.digital.read_bytes())
        with self.assertRaises(extract.SourceChanged):
            document.page_image(1)

    def test_A5_a_loaded_document_without_its_source_refuses_images_but_exports_text(self):
        document = self.read(self.scanned)
        path = self.tmp / "doc.extract.json"
        document.save(path)
        elsewhere = self.tmp / "elsewhere"
        elsewhere.mkdir()
        (elsewhere / path.name).write_bytes(path.read_bytes())
        loaded = extract.load(elsewhere / path.name)
        self.assertIsNone(loaded.source)
        with self.assertRaises(extract.SourceChanged):
            loaded.page_image(1)
        self.assertEqual(len(loaded.export(self.tmp / "text", ["txt", "json"])), 2)


class EventsTest(ApiTest):
    def test_A4_events_come_in_order_and_progress_only_grows(self):
        seen = []
        self.read(self.scanned, on_event=seen.append)
        names = [type(e).__name__ for e in seen]
        self.assertEqual(names[0], "Started")
        self.assertEqual(names[-2:], ["DocumentRead", "Finished"])
        self.assertEqual(seen[-1].status, "done")
        progress = [e.done for e in seen if isinstance(e, events.Progress)]
        self.assertEqual(progress, sorted(progress))
        self.assertEqual(progress[-1], 2)
        first_page_read = names.index("PageRead")
        self.assertLess(names.index("Started"), first_page_read)
        self.assertLess(first_page_read, names.index("DocumentRead"))

    def test_A4_events_count_the_doubtful_words_of_each_page_and_of_the_document(self):
        seen = []
        self.read(self.scanned, DoubtingGuard({"word3": None}), on_event=seen.append)
        self.assertEqual([e.doubtful for e in seen if isinstance(e, events.PageRead)], [1, 1])
        self.assertEqual(next(e for e in seen if isinstance(e, events.DocumentRead)).doubtful, 2)

    def test_A4_events_serialise_to_plain_dictionaries(self):
        seen = []
        self.read(self.digital, on_event=seen.append)
        json.dumps([e.to_dict() for e in seen])
        self.assertEqual(seen[0].to_dict(), {"event": "Started", "documents": 1, "pages": 2})

    def test_A4_cancel_stops_within_two_seconds_raises_and_leaves_no_output(self):
        scan = text_page_image(SPANISH_LINES[:3], SANS, 24, 827, 1169)
        path = scanned_pdf(self.tmp / "long.pdf", [scan] * 40)
        token = extract.CancelToken()
        seen = []
        stopped = {}

        def handler(event):
            seen.append(event)
            if isinstance(event, events.PageRead) and "at" not in stopped:
                stopped["at"] = time.perf_counter()
                token.cancel()

        with mock.patch.dict(os.environ, {"EXTRACT_FAKE_SECONDS": "0.3"}):
            reader = self.reader()
            with self.assertRaises(extract.Cancelled):
                reader.read(path, on_event=handler, cancel=token)
            self.assertLess(time.perf_counter() - stopped["at"], 2.0)
        self.assertEqual(seen[-1], events.Finished("stopped"))
        self.assertFalse((self.tmp / "out").exists())
        self.assertEqual([p for p in (self.tmp / "work").rglob("*") if p.suffix in (".part", ".partial")], [])


class ErrorsTest(ApiTest):
    def test_A5_each_error_comes_from_its_case_with_its_code(self):
        cases = []
        not_pdf = self.tmp / "not.pdf"
        not_pdf.write_bytes(b"plain text")
        cases.append((extract.UnsupportedFile, "error.unsupported_file", lambda: self.read(not_pdf)))
        locked = self.tmp / "locked.pdf"
        with pikepdf.Pdf.new() as pdf:
            pdf.add_blank_page()
            pdf.save(locked, encryption=pikepdf.Encryption(owner="o", user="u"))
        cases.append((extract.PasswordProtected, "error.password_protected", lambda: self.read(locked)))
        document = self.read(self.scanned)
        cases.append((extract.OutputExists, "error.output_exists", lambda: document.export(self.scanned)))
        cases.append((extract.SourceChanged, "error.source_changed", lambda: self._export_from_other_source()))
        token = extract.CancelToken()
        token.cancel()
        cases.append((extract.Cancelled, "error.cancelled", lambda: self.read(self.scanned, cancel=token)))
        for error, code, case in cases:
            with self.subTest(code=code), self.assertRaises(error) as caught:
                case()
            self.assertEqual(caught.exception.code, code)
            self.assertIsInstance(caught.exception, extract.ExtractError)

    def _export_from_other_source(self):
        document = self.read(self.scanned)
        document.source.write_bytes(self.digital.read_bytes())
        document.export(self.tmp / "o", ["pdf"])

    def test_A5_an_engine_that_cannot_start_is_an_engine_load_error(self):
        with mock.patch.object(Reader, "_engine_factory", "tests.fake_engine:does_not_exist"):
            with self.assertRaises(extract.EngineLoadError) as caught:
                self.read(self.scanned)
        self.assertEqual(caught.exception.code, "error.engine_load")


class WorkersTest(ApiTest):
    def test_A8_a_batch_of_20_documents_loads_the_workers_once_and_close_frees_them(self):
        import psutil

        paths = [
            scanned_pdf(self.tmp / f"d{n}.pdf", [text_page_image(SPANISH_LINES[:3], SANS, 20, 827, 1169)])
            for n in range(20)
        ]
        reader = self.reader()
        for path in paths:
            reader.read(path)
        started = [e for e in reader.job_log if e["code"] == "worker.started"]
        self.assertEqual(len(started), 1)

        def workers():
            return [c for c in psutil.Process().children(recursive=True) if "spawn_main" in " ".join(c.cmdline())]

        children = workers()
        self.assertEqual(len(children), 1)
        reader.close()
        time.sleep(0.5)
        self.assertEqual(workers(), [])


class UiImportsTest(unittest.TestCase):
    ALLOWED = "extract"

    @classmethod
    def violations(cls, source: str) -> list[str]:
        found = []
        for node in ast.walk(ast.parse(source)):
            if isinstance(node, ast.Import):
                names = [a.name for a in node.names]
            elif isinstance(node, ast.ImportFrom):
                if node.level == 1:
                    continue  # a sibling module of the interface itself
                if node.level > 1:
                    found.append("a relative import out of the interface")
                    continue
                names = [node.module or ""]
            else:
                continue
            for name in names:
                root = name.split(".")[0]
                if root not in sys.stdlib_module_names and root != cls.ALLOWED:
                    found.append(root)
                if root == cls.ALLOWED and "." in name:
                    found.append(name)
        return found

    def test_A6_the_scanner_catches_what_the_interface_must_not_import(self):
        self.assertEqual(
            self.violations("import json\nfrom http import server\nimport extract\nfrom extract import read"), []
        )
        self.assertEqual(self.violations("import numpy"), ["numpy"])
        self.assertEqual(self.violations("from extract.guard import Guard"), ["extract.guard"])
        self.assertEqual(self.violations("from . import x\nfrom .server import serve"), [])
        self.assertEqual(self.violations("from .. import engine"), ["a relative import out of the interface"])
        self.assertEqual(self.violations("from ..guard import Guard"), ["a relative import out of the interface"])

    def test_A6_the_interface_imports_only_extract_and_the_standard_library(self):
        folder = Path(extract.__file__).parent / "ui"
        for path in sorted(folder.rglob("*.py")) if folder.is_dir() else []:
            with self.subTest(path=path.name):
                self.assertEqual(self.violations(path.read_text(encoding="utf-8")), [])


if __name__ == "__main__":
    unittest.main()


class FixesApiTest(ApiTest):
    """A read with the guard's proposals: the mode on `read`, the events and the document's own mode switch."""

    def setUp(self):
        super().setUp()
        names = [
            "alfa",
            "beta",
            "gamma",
            "delta",
            "epsilon",
            "zeta",
            "eta",
            "theta",
            "iota",
            "kappa",
            "lambda",
            "sigma",
        ]
        words = [[f"{names[i % 12]}{'x' * (i // 12)}", 0.9, [10 * i, 0, 10 * i + 8, 10]] for i in range(60)]
        page = self.tmp / "words.json"
        page.write_text(json.dumps([words[:30], words[30:]]))
        patcher = mock.patch.dict(os.environ, {"EXTRACT_FAKE_PAGE": str(page)})
        patcher.start()
        self.addCleanup(patcher.stop)

    def guard(self):
        return FixGuard(
            {
                "p1w3": fix("text", "accent", "tercera", ("p1w3",)),
                "p1w5": fix("text", "case", "Quinta", ("p1w5",), modes=("more",)),
            }
        )

    def test_C6_read_applies_the_proposals_of_the_mode_and_reports_them(self):
        seen = []
        document = self.read(self.scanned, self.guard(), on_event=seen.append, correct="safe")
        self.assertEqual((document.mode, document.fixes["applied"], document.word("p1w3").text), ("safe", 1, "tercera"))
        self.assertEqual(document.word("p1w5").status, "doubtful")
        self.assertEqual(sum(e.fixed for e in seen if type(e).__name__ == "PageRead"), 1)
        self.assertEqual([e.fixed for e in seen if type(e).__name__ == "DocumentRead"], [1])
        document.set_mode("more")
        self.assertEqual(document.fixes["applied"], 2)

    def test_C6_the_option_sets_the_default_mode_and_off_is_the_plain_read(self):
        plain = self.read(self.scanned, FakeGuard())
        with_options = Reader(
            extract.Options(work_dir=self.tmp / "work", processes=1, threads=1, correct="more"),
            guard=self.guard(),
            lexicon=LEXICON,
        )
        self.addCleanup(with_options.close)
        self.assertEqual(with_options.read(self.scanned).fixes["applied"], 2)
        off = self.read(self.scanned, self.guard(), correct="off")
        self.assertEqual([w.text for w in off.words()], [w.text for w in plain.words()])
        self.assertEqual(off.fixes["applied"], 0)

    def test_C6_a_wrong_mode_is_refused(self):
        with self.assertRaises(ValueError):
            self.read(self.scanned, self.guard(), correct="maybe")
        document = self.read(self.scanned, self.guard())
        with self.assertRaises(ValueError):
            document.set_mode("maybe")

    def test_C6_the_new_names_are_public(self):
        for name in ("Fix", "Reading"):
            self.assertIn(name, extract.__all__)
