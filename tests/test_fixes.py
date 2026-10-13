"""Applying the guard's proposals: modes, undo, structure fixes, token view, page brake, digit fixes, and the
JSON, text, PDF and Excel that carry them (C1-C5, C9, C10)."""

import json
import tempfile
import unittest
from dataclasses import replace
from pathlib import Path

import jsonschema
import openpyxl
import pypdfium2 as pdfium

from extract import fixes
from extract.api import Document
from extract.events import Warning as JobWarning
from extract.outputs import export, jsondoc
from extract.outputs.text import to_text
from extract.types import DigitSuggestion, DocumentData, NumberResult, base_id, page_of
from tests.checks_oracle import DNI
from tests.fake_fix_guard import fix, with_fixes
from tests.support import SANS, SPANISH_LINES, grid_reading, scanned_pdf, text_page_image

SCHEMA = json.loads((Path(jsondoc.__file__).parent.parent / "schema" / "document.schema.json").read_text())

# 120 words: two fixes touch 5 of them at most, well under the 8 % brake.
WORDS = [f"palabra{chr(97 + i % 26)}{chr(97 + (i // 26) % 26)}" for i in range(120)]
WORDS[3], WORDS[4], WORDS[5] = "fac", "tura", "de"  # a join of three
WORDS[10], WORDS[11] = "im", "porte"  # a join of two
WORDS[20] = "Totalfactura"  # a segment
WORDS[30] = "ı"  # a drop
WORDS[40] = "telefono"  # a text fix
WORDS[41] = "ES91-2100"  # a protected word (digits)
WORDS[42] = "juan@mail"  # a protected word (@)
WORDS[43] = "Telefono"  # a plain title-case word: allowed
WORDS[44] = "iPhoneXR"  # two inner capitals: protected


def page(texts=None):
    return grid_reading(1, list(texts or WORDS), per_row=8)


def ids(*numbers):
    return tuple(f"p1w{n}" for n in numbers)


def proposals():
    return {
        "p1w3": fix("join", "structure", "factura", ids(3, 4, 5), modes=("safe", "more")),
        "p1w10": fix("join", "structure", "importe", ids(10, 11), modes=("more",)),
        "p1w20": fix("segment", "structure", "Total factura", ids(20), modes=("more",)),
        "p1w30": fix("drop", "structure", "", ids(30), modes=("more",)),
        "p1w40": fix("text", "accent", "teléfono", ids(40)),
    }


def texts(p):
    return [w.text for w in p.words]


class ModeTest(unittest.TestCase):
    def setUp(self):
        self.base = with_fixes(page(), proposals())

    def test_C1_off_applies_nothing_and_shows_the_suggestion_of_a_doubtful_word(self):
        applied, count = fixes.apply_fixes(self.base, "off")
        self.assertEqual(count, 0)
        self.assertEqual(texts(applied), WORDS)
        word = applied.words[40]
        self.assertEqual((word.status, word.suggestion, word.suggestion_conf), ("doubtful", "teléfono", 0.97))
        self.assertIsNone(applied.words[3].suggestion)  # a structure fix is not a text suggestion

    def test_C1_safe_applies_the_fixes_whose_modes_include_it(self):
        applied, count = fixes.apply_fixes(self.base, "safe")
        self.assertEqual(count, 2)
        self.assertEqual(applied.words[40].text, "teléfono")
        self.assertEqual(applied.words[3].text, "factura")
        self.assertEqual(applied.words[10].text, "im")  # a more-only fix
        self.assertEqual(applied.words[20].text, "Totalfactura")

    def test_C1_more_applies_all_of_them(self):
        applied, count = fixes.apply_fixes(self.base, "more")
        self.assertEqual(count, 5)
        self.assertEqual(applied.words[10].text, "importe")
        self.assertEqual(applied.words[20].text, "Total factura")
        self.assertEqual(applied.words[30].status, "absorbed")

    def test_A2_3_an_applied_text_fix_sets_the_fields_of_a_fixed_word(self):
        word = fixes.apply_fixes(self.base, "safe")[0].words[40]
        self.assertEqual(
            (word.status, word.text, word.original, word.corrected_by, word.suggestion, word.guard_status),
            ("fixed", "teléfono", "telefono", "guard", None, "doubtful"),
        )
        self.assertEqual((word.p_wrong, word.reasons), (0.7, self.base.words[40].reasons))

    def test_C1_a_fix_on_an_ok_word_applies_too(self):
        ok = with_fixes(page(), {"p1w40": proposals()["p1w40"]}, status="ok")
        self.assertEqual(fixes.apply_fixes(ok, "safe")[0].words[40].status, "fixed")

    def test_C1_a_word_the_user_decided_never_moves(self):
        decided = replace(
            self.base, words=tuple(replace(w, review="accepted") if w.id == "p1w40" else w for w in self.base.words)
        )
        for mode in ("off", "safe", "more"):
            self.assertEqual(fixes.apply_fixes(decided, mode)[0].words[40].text, "telefono")
        applied = fixes.apply_fixes(self.base, "safe")[0]
        kept = replace(
            applied, words=tuple(replace(w, review="accepted") if w.id == "p1w40" else w for w in applied.words)
        )
        self.assertEqual(fixes.apply_fixes(kept, "off")[0].words[40].text, "teléfono")  # a kept fix stays

    def test_C2_going_back_to_off_gives_the_off_page_and_set_mode_equals_reading_with_that_mode(self):
        off = fixes.apply_fixes(self.base, "off")[0]
        for mode in ("safe", "more"):
            direct = fixes.apply_fixes(self.base, mode)[0]
            via = fixes.apply_fixes(fixes.apply_fixes(fixes.apply_fixes(self.base, "more")[0], "off")[0], mode)[0]
            self.assertEqual(via, direct)
            self.assertEqual(fixes.apply_fixes(direct, "off")[0], off)

    def test_C2_undo_all_gives_the_off_text_ids_boxes_and_statuses(self):
        off = fixes.apply_fixes(self.base, "off")[0]
        applied = fixes.apply_fixes(self.base, "more")[0]
        undone, count = fixes.undo_words(applied)
        self.assertEqual(count, 5)
        self.assertEqual(
            [(w.id, w.text, w.box, w.status, w.suggestion) for w in undone.words],
            [(w.id, w.text, w.box, w.status, w.suggestion) for w in off.words],
        )
        # and a mode switch never applies them again
        self.assertEqual(texts(fixes.apply_fixes(undone, "more")[0]), WORDS)

    def test_C1_property_a_fix_is_applied_only_if_the_mode_is_in_its_modes(self):
        for modes in ((), ("safe",), ("more",), ("safe", "more")):
            one = with_fixes(page(), {"p1w40": fix("text", "accent", "teléfono", ids(40), modes=modes)})
            for mode in ("off", "safe", "more"):
                done = fixes.apply_fixes(one, mode)[0].words[40].status == "fixed"
                self.assertEqual(done, mode in modes, (modes, mode))

    def test_C1_overlapping_fixes_are_applied_in_reading_order_and_the_later_is_skipped(self):
        overlap = with_fixes(
            page(),
            {
                "p1w3": fix("join", "structure", "factura", ids(3, 4, 5)),
                "p1w4": fix("text", "accent", "TURA", ids(4)),
                "p1w5": fix("join", "structure", "dex", ids(5, 6)),
            },
        )
        applied = fixes.apply_fixes(overlap, "safe")[0]
        self.assertEqual(applied.words[3].text, "factura")
        self.assertEqual(applied.words[4].status, "absorbed")
        self.assertEqual(applied.words[5].status, "absorbed")
        self.assertEqual(applied.words[6].text, WORDS[6])


class StructureTest(unittest.TestCase):
    def setUp(self):
        self.base = with_fixes(page(), proposals())
        self.applied = fixes.apply_fixes(self.base, "more")[0]

    def test_C9_join_of_three_takes_the_union_box_and_absorbs_the_others(self):
        first = self.applied.words[3]
        boxes = [w.box for w in self.base.words[3:6]]
        self.assertEqual(
            first.box,
            (min(b[0] for b in boxes), min(b[1] for b in boxes), max(b[2] for b in boxes), max(b[3] for b in boxes)),
        )
        for member in self.applied.words[4:6]:
            self.assertEqual((member.status, member.text, member.absorbed_by), ("absorbed", "", "p1w3"))
        self.assertEqual([m.original for m in self.applied.words[4:6]], ["tura", "de"])

    def test_C9_join_of_two(self):
        self.assertEqual((self.applied.words[10].text, self.applied.words[11].status), ("importe", "absorbed"))

    def test_C9_the_token_view_skips_absorbed_words_and_splits_a_segment_in_proportion_to_the_font_widths(self):
        tokens = fixes.tokens(self.applied)
        by_id = {t.id: t for t in tokens}
        self.assertNotIn("p1w4", by_id)
        self.assertNotIn("p1w30", by_id)
        a, b = by_id["p1w20#1"], by_id["p1w20#2"]
        self.assertEqual((a.text, b.text), ("Total", "factura"))
        word = self.applied.words[20]
        self.assertAlmostEqual(a.box[0], word.box[0])
        self.assertAlmostEqual(b.box[2], word.box[2])
        self.assertLess(a.box[2], b.box[0])
        self.assertAlmostEqual(
            (a.box[2] - a.box[0]) / (b.box[2] - b.box[0]),
            fixes._helvetica_width("Total") / fixes._helvetica_width("factura"),
        )
        self.assertEqual(base_id("p1w20#2"), "p1w20")
        self.assertEqual(page_of("p1w20#2"), 1)

    def test_C9_a_segment_keeps_one_word_and_its_id_and_box(self):
        word = self.applied.words[20]
        self.assertEqual((word.id, word.box, word.text), ("p1w20", self.base.words[20].box, "Total factura"))

    def test_C9_drop_has_no_owner_and_undo_brings_the_word_back(self):
        word = self.applied.words[30]
        self.assertEqual((word.status, word.text, word.absorbed_by, word.original), ("absorbed", "", None, "ı"))
        undone = fixes.undo_words(self.applied, "p1w30")[0]
        self.assertEqual((undone.words[30].text, undone.words[30].status), ("ı", "doubtful"))

    def test_C9_undo_of_a_member_undoes_the_whole_join_and_restores_boxes_exactly(self):
        undone, count = fixes.undo_words(self.applied, "p1w4")
        self.assertEqual(count, 1)
        for i in (3, 4, 5):
            self.assertEqual(
                (undone.words[i].text, undone.words[i].box, undone.words[i].status),
                (self.base.words[i].text, self.base.words[i].box, "doubtful" if i == 3 else "ok"),
            )
        self.assertEqual(undone.words[3].review, "accepted")

    def test_C9_protected_words_never_receive_a_fix(self):
        bad = {
            "p1w41": fix("text", "letters", "ES91-2I00", ids(41)),
            "p1w42": fix("text", "letters", "juan@maiI", ids(42)),
            "p1w44": fix("text", "case", "iPhonexr", ids(44)),
            "p1w43": fix("text", "accent", "Teléfono", ids(43)),
        }
        applied = fixes.apply_fixes(with_fixes(page(), bad), "more")[0]
        self.assertEqual([w.status for w in applied.words[41:45]], ["doubtful", "doubtful", "fixed", "doubtful"])
        for text in (
            "a1",
            "a@b",
            "a/b",
            "a_b",
            "iPhoneXR",
            "ABc De",
            "GrafCan",
            "McDonald",
            "a#b",
            "a\\b",
            "a:b",
            "http:",
            "obra-",
            "l'",
            "x" * 41,
        ):
            self.assertTrue(fixes.is_protected(text), text)
        for text in ("Total", "ñandú", "TOTAL", "Urbanísticos:", "x"):
            self.assertFalse(fixes.is_protected(text), text)

    def test_C9_a_join_needs_consecutive_words_of_one_line(self):
        far = with_fixes(page(), {"p1w3": fix("join", "structure", "xx", ids(3, 5))})
        self.assertEqual(fixes.apply_fixes(far, "more")[1], 0)
        other_line = with_fixes(page(), {"p1w7": fix("join", "structure", "xx", ids(7, 8))})  # row break after 8 words
        self.assertEqual(fixes.apply_fixes(other_line, "more")[1], 0)


class DropOfMarksTest(unittest.TestCase):
    """The guard drops specks of marks (`·`, `|`, `.`, `•`); the app must apply them (a regression found in review)."""

    def drop(self, text):
        texts_ = list(WORDS)
        texts_[50] = text
        return fixes.apply_fixes(with_fixes(page(texts_), {"p1w50": fix("drop", "structure", "", ids(50))}), "more")[1]

    def test_C9_a_drop_of_a_speck_of_marks_is_applied(self):
        for text in ("·", "|", ".", "•", ",", "ı", "l", "o"):
            with self.subTest(text=text):
                self.assertEqual(self.drop(text), 1)

    def test_C9_a_drop_of_a_hyphen_apostrophe_digit_or_protected_character_is_not(self):
        for text in ("-", "'", "a1", "a/b", "@", "_"):
            with self.subTest(text=text):
                self.assertEqual(self.drop(text), 0)

    def test_D4_op_counts_a_token_of_marks_is_protected_for_every_op_but_drop(self):
        for text in ("·", "|", ".", "•", ","):
            with self.subTest(text=text):
                self.assertFalse(fixes.is_protected(text, "drop"))
                for op in ("text", "join", "segment"):
                    self.assertTrue(fixes.is_protected(text, op), op)
                    texts_ = list(WORDS)
                    texts_[50] = text
                    group = ids(50, 51) if op == "join" else ids(50)
                    proposed = with_fixes(page(texts_), {"p1w50": fix(op, "structure", "ab cd", group)})
                    self.assertEqual(fixes.apply_fixes(proposed, "more")[1], 0, op)

    def test_C9_property_a_speck_of_marks_is_dropped_unless_it_holds_a_protected_character_or_a_dash(self):
        for text in '·|.,•:;!?*~+=<>()[]{}"':
            with self.subTest(text=text):
                self.assertEqual(self.drop(text), 1)
        for text in "-'\u2014\u2013/@_#\\":
            with self.subTest(text=text):
                self.assertEqual(self.drop(text), 0)


class ProtectedPropertyTest(unittest.TestCase):
    def test_C9_property_a_token_with_any_protected_trait_never_receives_a_fix(self):
        traits = ["{}1", "{}@x", "{}/x", "{}_x", "{}#x", "{}\\x", "{}:x", "{}-", "{}'", "{}xAB", "{}xCD", "http:{}"]
        for trait in traits:
            text = trait.format("abc")
            with self.subTest(text=text):
                self.assertTrue(fixes.is_protected(text))
                for op in ("text", "segment", "drop", "join"):
                    group = ids(50, 51) if op == "join" else ids(50)
                    texts_ = list(WORDS)
                    texts_[50] = text
                    proposed = with_fixes(page(texts_), {"p1w50": fix(op, "structure", "ab cd", group)})
                    self.assertEqual(fixes.apply_fixes(proposed, "more")[1], 0, (op, text))


class BrakeTest(unittest.TestCase):
    def test_C10_the_edge_is_exactly_eight_percent_and_a_join_counts_its_words(self):
        hundred = WORDS[:100]
        eight = {f"p1w{i}": fix("text", "accent", f"x{i}", ids(i)) for i in range(60, 68)}
        self.assertEqual(fixes.apply_fixes(with_fixes(page(hundred), eight), "safe")[1], 8)
        nine = {f"p1w{i}": fix("text", "accent", f"x{i}", ids(i)) for i in range(60, 69)}
        self.assertEqual(fixes.apply_fixes(with_fixes(page(hundred), nine), "safe")[1], 0)
        joins = {f"p1w{i}": fix("join", "structure", "jj", ids(i, i + 1, i + 2)) for i in (56, 64, 72)}  # 9 words
        self.assertEqual(fixes.apply_fixes(with_fixes(page(hundred), joins), "safe")[1], 0)
        two = {f"p1w{i}": fix("join", "structure", "jj", ids(i, i + 1, i + 2)) for i in (56, 64)}  # 6 words
        self.assertEqual(fixes.apply_fixes(with_fixes(page(hundred), two), "safe")[1], 2)

    def test_C10_over_eight_percent_of_the_page_applies_nothing_and_warns(self):
        many = {f"p1w{i}": fix("text", "accent", f"x{i}", ids(i)) for i in range(0, 12)}  # 10 % of 120
        events = []
        applied, count = fixes.apply_fixes(with_fixes(page(), many), "safe", events.append)
        self.assertEqual(count, 0)
        self.assertEqual(texts(applied), WORDS)
        self.assertEqual(len(events), 1)
        self.assertIsInstance(events[0], JobWarning)
        self.assertEqual(events[0].code, "guard.warning.fixes_withheld")
        self.assertEqual(events[0].params["page"], 1)
        self.assertAlmostEqual(events[0].params["share"], 0.1)
        self.assertEqual(applied.words[0].suggestion, "x0")  # the proposals stay as suggestions

    def test_C10_up_to_eight_percent_is_applied(self):
        few = {f"p1w{i}": fix("text", "accent", f"x{i}", ids(i)) for i in range(0, 9)}  # 7.5 %
        events = []
        self.assertEqual(fixes.apply_fixes(with_fixes(page(), few), "safe", events.append)[1], 9)
        self.assertEqual(events, [])

    def test_C10_off_never_warns(self):
        many = {f"p1w{i}": fix("text", "accent", f"x{i}", ids(i)) for i in range(0, 30)}
        events = []
        fixes.apply_fixes(with_fixes(page(), many), "off", events.append)
        self.assertEqual(events, [])


def number(word_id, text, position, old, new, verified=True, rule="checks.rule.dni_letter"):
    return NumberResult(
        (word_id,), text, None, "dni", "mismatch", (), DigitSuggestion(position, old, new, verified, rule)
    )


class DigitTest(unittest.TestCase):
    def setUp(self):
        texts_ = list(WORDS)
        texts_[50] = DNI
        self.page = replace(
            page(texts_),
            words=tuple(
                replace(
                    w, char_alts=tuple(((c, 0.9),) + ((("9", 0.05),) if i == 4 else ()) for i, c in enumerate(w.text))
                )
                if w.id == "p1w50"
                else w
                for w in page(texts_).words
            ),
        )

    def run_digit(self, suggestion, mode="safe", the_page=None):
        return fixes.apply_digit_fixes([the_page or self.page], [suggestion], mode)

    def test_C3_a_verified_digit_among_the_alternatives_is_fixed_by_the_rule(self):
        pages, count = self.run_digit(number("p1w50", DNI, 4, "5", "9"))
        word = pages[0].words[50]
        self.assertEqual(
            (count, word.text, word.status, word.corrected_by), (1, DNI[:4] + "9" + DNI[5:], "fixed", "guard")
        )
        self.assertEqual(
            (word.fix.source, word.fix.rule, word.original),
            ("rule:checks.rule.dni_letter", "checks.rule.dni_letter", DNI),
        )

    def test_C3_a_word_that_two_results_point_at_is_fixed_once_and_by_the_first(self):
        alts = tuple(((c, 0.9),) + ((("9", 0.05), ("3", 0.04)) if i == 4 else ()) for i, c in enumerate(DNI))
        the_page = replace(
            self.page, words=tuple(replace(w, char_alts=alts) if w.id == "p1w50" else w for w in self.page.words)
        )
        first, second = number("p1w50", DNI, 4, "5", "9"), number("p1w50", DNI, 4, "5", "3")
        pages, count = fixes.apply_digit_fixes([the_page], [first, second], "safe")
        self.assertEqual((count, pages[0].words[50].text), (1, DNI[:4] + "9" + DNI[5:]))

    def test_C3_a_token_id_of_a_segmented_word_is_never_fixed_even_if_a_word_has_that_id(self):
        the_page = replace(
            self.page, words=tuple(replace(w, id="p1w50#1") if w.id == "p1w50" else w for w in self.page.words)
        )
        pages, count = self.run_digit(number("p1w50#1", DNI, 4, "5", "9"), the_page=the_page)
        self.assertEqual((count, pages[0].words[50].text), (0, DNI))

    def test_B43_in_safe_arithmetic_is_a_suggestion_and_a_check_digit_is_applied(self):
        literal = {
            "checks.rule.line_amount",
            "checks.rule.lines_sum",
            "checks.rule.vat_amount",
            "checks.rule.withholding_amount",
            "checks.rule.total",
        }
        self.assertEqual(set(fixes.ARITHMETIC_RULES), literal)
        for rule in sorted(literal):
            suggestion = number("p1w50", DNI, 4, "5", "9", rule=rule)
            self.assertEqual(self.run_digit(suggestion, "safe")[1], 0, rule)
            self.assertEqual(self.run_digit(suggestion, "more")[1], 1, rule)
        for rule in (
            "checks.rule.dni_letter",
            "checks.rule.nie_letter",
            "checks.rule.iban_mod97",
            "checks.rule.cif_control",
        ):
            self.assertEqual(self.run_digit(number("p1w50", DNI, 4, "5", "9", rule=rule), "safe")[1], 1, rule)

    def test_B43_an_arithmetic_fix_made_in_more_is_marked_for_more_only_is_not_kept_in_safe_and_can_be_undone(self):
        suggestion = number("p1w50", DNI, 4, "5", "9", rule="checks.rule.total")
        made = self.run_digit(suggestion, "more")[0][0]
        fix_ = made.words[50].fix
        self.assertEqual(fix_.modes, ("more",))
        self.assertEqual(fixes.apply_fixes(made, "safe")[0].words[50].text, DNI)
        self.assertEqual(fixes.undo_words(made, "p1w50")[0].words[50].text, DNI)
        control = self.run_digit(number("p1w50", DNI, 4, "5", "9"), "safe")[0][0]
        self.assertEqual(control.words[50].fix.modes, ("safe", "more"))

    def test_B43_end_to_end_the_five_arithmetic_rules_of_checks_run_are_suggestions_in_safe_and_applied_in_more(self):
        from bench import inject_fixes
        from bench.corpus.sidecar_inputs import build, generated_truths
        from extract import checks

        found = {}
        for truth in generated_truths()[:3]:
            inputs = build(truth)
            for word_id in inputs.ids.values():
                if isinstance(word_id, tuple) or not any(w.id == word_id for w in inputs.words):
                    continue
                text = inputs.text_of(word_id)
                at = next((i for i, c in enumerate(text) if c.isdigit() and inject_fixes.alternatives(c)), None)
                if at is None:
                    continue
                wrong = inputs.with_texts(
                    {word_id: text[:at] + inject_fixes.alternatives(text[at])[0] + text[at + 1 :]}
                )
                results = checks.run(wrong.words, wrong.fields, truth["language"])
                for r in results:
                    if r.suggestion is not None and r.suggestion.verified and len(r.word_ids) == 1:
                        found.setdefault(r.suggestion.rule, (wrong, [r], inputs))
        self.assertTrue(set(fixes.ARITHMETIC_RULES) <= set(found), sorted(found))
        for rule, (wrong, results, truth_inputs) in found.items():
            alone = inject_fixes.apply_verified(wrong, results, True, truth=truth_inputs, mode="safe")[1]
            more = inject_fixes.apply_verified(wrong, results, True, truth=truth_inputs, mode="more")[1]
            self.assertEqual(len(alone), 0 if rule in fixes.ARITHMETIC_RULES else 1, rule)
            self.assertEqual(len(more), 1, rule)

    def test_C3_a_digit_that_is_not_among_the_alternatives_is_not_changed(self):
        self.assertEqual(self.run_digit(number("p1w50", DNI, 4, "5", "8"))[1], 0)

    def test_C3_an_unverified_suggestion_is_not_applied(self):
        self.assertEqual(self.run_digit(number("p1w50", DNI, 4, "5", "9", verified=False))[1], 0)

    def test_C3_off_never_changes_a_digit(self):
        self.assertEqual(self.run_digit(number("p1w50", DNI, 4, "5", "9"), "off")[1], 0)

    def test_C3_a_word_without_alternatives_or_decided_by_the_user_is_not_changed(self):
        bare = replace(self.page, words=tuple(replace(w, char_alts=None) for w in self.page.words))
        self.assertEqual(self.run_digit(number("p1w50", DNI, 4, "5", "9"), the_page=bare)[1], 0)
        decided = replace(
            self.page, words=tuple(replace(w, review="accepted") if w.id == "p1w50" else w for w in self.page.words)
        )
        self.assertEqual(self.run_digit(number("p1w50", DNI, 4, "5", "9"), the_page=decided)[1], 0)

    def test_C3_the_guards_own_proposals_never_touch_a_word_with_a_digit(self):
        proposal = with_fixes(
            self.page, {"p1w50": fix("text", "letters", DNI[:-1] + "S", ids(50), modes=("safe", "more"))}
        )
        self.assertEqual(fixes.apply_fixes(proposal, "more")[0].words[50].text, DNI)

    def test_C3_the_digit_fix_can_be_undone_and_a_mode_switch_forgets_and_remakes_it(self):
        pages, _ = self.run_digit(number("p1w50", DNI, 4, "5", "9"))
        undone = fixes.undo_words(pages[0], "p1w50")[0].words[50]
        self.assertEqual((undone.text, undone.status, undone.review), (DNI, "unscored", "accepted"))
        off = fixes.apply_fixes(pages[0], "off")[0].words[50]
        self.assertEqual((off.text, off.status), (DNI, "unscored"))


def document(source=None, proposals_=None, mode="off", status="doubtful", texts_=None):
    base = with_fixes(page(texts_), proposals_ if proposals_ is not None else proposals(), status)
    applied = fixes.apply_fixes(base, mode)[0]
    data = DocumentData(
        source=source.name if source else "doc.pdf",
        sha256="0" * 64,
        pages=(applied,),
        engine="fake",
        guard="fake-guard",
        options={"languages": ["es", "en"], "review": "balanced", "correct": mode, "second_reading": False},
        doc_type="generic",
        locale="es",
    )
    return Document(data, source)


class DocumentTest(unittest.TestCase):
    def test_C2_set_mode_re_applies_and_off_brings_the_read_back(self):
        doc = document()
        off_text = [w.text for w in doc.words()]
        doc.set_mode("more")
        self.assertEqual((doc.mode, doc.fixes["applied"]), ("more", 5))
        doc.set_mode("safe")
        self.assertEqual(doc.fixes["applied"], 2)
        doc.set_mode("off")
        self.assertEqual([w.text for w in doc.words()], off_text)
        self.assertEqual(doc.fixes["applied"], 0)

    def test_C2_set_mode_equals_reading_with_that_mode(self):
        for mode in ("safe", "more"):
            doc = document()
            doc.set_mode(mode)
            self.assertEqual(doc.pages, document(mode=mode).pages)

    def test_C2_undo_and_undo_page_and_undo_all(self):
        doc = document(mode="more")
        doc.undo("p1w40")
        self.assertEqual(doc.word("p1w40").text, "telefono")
        self.assertEqual((doc.fixes["applied"], doc.fixes["undone"]), (4, 1))
        with self.assertRaises(ValueError):
            doc.undo("p1w40")
        self.assertEqual(doc.undo_page(1), 4)
        self.assertEqual(doc.fixes["applied"], 0)
        doc.set_mode("more")  # an undone fix is not applied again
        self.assertEqual(doc.fixes["applied"], 0)
        self.assertEqual([w.text for w in doc.words()], WORDS)
        self.assertEqual(doc.undo_all(), 0)

    def test_C2_undo_all_gives_the_off_text_ids_boxes_and_statuses(self):
        off = document()
        doc = document(mode="more")
        doc.undo_all()
        self.assertEqual(
            [(w.id, w.text, w.box, w.status) for w in doc.words()],
            [(w.id, w.text, w.box, w.status) for w in off.words()],
        )

    def test_C2_a_user_correction_is_not_moved_by_a_mode_switch(self):
        doc = document()
        doc.correct("p1w40", "teléfono!")
        doc.set_mode("more")
        self.assertEqual(doc.word("p1w40").text, "teléfono!")

    def test_C2_the_set_of_word_ids_is_the_same_in_every_mode(self):
        doc = document()
        before = [w.id for w in doc.words()]
        for mode in ("more", "safe", "off"):
            doc.set_mode(mode)
            self.assertEqual([w.id for w in doc.words()], before)

    def test_C9_tokens_of_a_document_page(self):
        doc = document(mode="more")
        texts_ = [t.text for t in doc.tokens(1)]
        self.assertIn("Total", texts_)
        self.assertIn("factura", texts_)
        self.assertEqual(len(texts_), 120 - 2 - 1 - 1 + 1)  # join3 -2, join2 -1, drop -1, segment +1

    def test_C10_the_page_brake_is_recorded_in_the_document_and_the_json(self):
        many = {f"p1w{i}": fix("text", "accent", f"x{i}", ids(i)) for i in range(0, 12)}
        doc = document(proposals_=many)
        doc.set_mode("safe")
        self.assertEqual(doc.fixes["applied"], 0)
        self.assertEqual(doc.fixes["withheld"][0]["page"], 1)
        raw = jsondoc.to_dict(doc._data)
        self.assertEqual(raw["options"]["fixes_withheld"][0]["page"], 1)
        jsonschema.validate(raw, SCHEMA)


class OutputsTest(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        pdf_page = text_page_image(SPANISH_LINES, SANS, 24, 1240, 1754)
        self.source = scanned_pdf(self.tmp / "doc.pdf", [pdf_page])
        self.doc = document(self.source, mode="more")
        from extract.engine.cache import file_sha256

        self.doc._data = replace(self.doc._data, sha256=file_sha256(self.source))

    def test_C4_the_json_v2_validates_for_fixed_absorbed_segmented_undone_and_doubtful_words(self):
        self.doc.undo("p1w40")
        raw = jsondoc.to_dict(self.doc._data)
        jsonschema.validate(raw, SCHEMA)
        self.assertEqual(raw["schema_version"], 2)
        by_id = {w["id"]: w for w in raw["pages"][0]["words"]}
        self.assertEqual(
            (by_id["p1w3"]["status"], by_id["p1w3"]["fix"]["applied"], by_id["p1w3"]["fix"]["op"]),
            ("fixed", True, "join"),
        )
        self.assertEqual((by_id["p1w4"]["status"], by_id["p1w4"]["absorbed_by"]), ("absorbed", "p1w3"))
        self.assertEqual((by_id["p1w20"]["text"], by_id["p1w20"]["fix"]["op"]), ("Total factura", "segment"))
        self.assertEqual((by_id["p1w40"]["status"], by_id["p1w40"]["fix"]["applied"]), ("doubtful", False))
        self.assertEqual((by_id["p1w0"]["status"], by_id["p1w0"]["fix"]), ("ok", None))

    def test_C4_the_fixes_block_adds_up_and_a_round_trip_is_equal(self):
        raw = jsondoc.to_dict(self.doc._data)
        block = raw["fixes"]
        self.assertEqual(block["mode"], "more")
        self.assertEqual(sum(block["by_source"].values()), block["applied"])
        self.assertEqual(sum(block["by_op"].values()), block["applied"])
        self.assertEqual(block["by_op"], {"drop": 1, "join": 2, "segment": 1, "text": 1})
        path = self.tmp / "d.extract.json"
        jsondoc.save(self.doc._data, path)
        self.assertEqual(jsondoc.load(path), jsondoc.exported(self.doc._data))

    def test_C4_a_loaded_document_can_be_undone_and_switched(self):
        path = self.tmp / "d.extract.json"
        self.doc.save(path)
        from extract.api import load

        loaded = load(path, self.source)
        loaded.undo("p1w3")
        self.assertEqual([loaded.word(f"p1w{i}").text for i in (3, 4, 5)], ["fac", "tura", "de"])
        self.assertEqual(loaded.word("p1w3").box, document().word("p1w3").box)
        loaded.set_mode("off")
        self.assertEqual(loaded.fixes["applied"], 0)

    def test_C5_text_pdf_and_excel_agree_and_absorbed_words_are_absent(self):
        paths = {p.suffix if p.suffix != ".pdf" else ".pdf": p for p in self.doc.export(self.tmp / "out")}
        text = paths[".txt"].read_text(encoding="utf-8").split()
        expected = [t.text for t in self.doc.tokens(1)]
        self.assertEqual(" ".join(text).split(), expected_split(expected))
        self.assertEqual(to_text(self.doc._data).split(), expected_split(expected))
        document_ = pdfium.PdfDocument(str(paths[".pdf"]))
        textpage = document_[0].get_textpage()
        layer = textpage.get_text_range().split()
        textpage.close()
        self.assertEqual(layer, expected_split([_ascii(t) for t in expected]))
        document_.close()
        book = openpyxl.load_workbook(paths[".xlsx"])
        sheet = book["Corrections"]
        rows = list(sheet.iter_rows(min_row=2, values_only=True))
        self.assertEqual(len(rows), self.doc.fixes["applied"])
        self.assertEqual(sorted(r[1] for r in rows), ["drop", "join", "join", "segment", "text"])
        self.assertTrue(all(r[8] == "applied" for r in rows))

    def test_C5_an_undone_fix_is_listed_as_undone_and_not_in_review(self):
        self.doc.undo("p1w40")
        from extract.outputs.excel import build_workbook

        book = build_workbook(self.doc._data)
        rows = {r[3]: r for r in book["Corrections"].iter_rows(min_row=2, values_only=True)}
        self.assertEqual(rows["telefono"][8], "undone")
        self.assertEqual(rows["telefono"][4], "telefono")
        reviewed = [r[1] for r in book["Review"].iter_rows(min_row=2, values_only=True)]
        self.assertNotIn("teléfono", reviewed)

    def test_C5_the_export_checks_still_pass_with_fixes(self):
        paths = export.export(self.doc._data, self.source, self.tmp / "all", export.FORMATS)
        self.assertEqual(len(paths), 4)


def expected_split(items):
    return [part for item in items for part in item.split()]


def _ascii(text):
    return text.encode("cp1252", "replace").decode("cp1252")


class MalformedFixTest(unittest.TestCase):
    def applied(self, proposal, owner="p1w5"):
        return fixes.apply_fixes(with_fixes(page(), {owner: proposal}), "more")[1]

    def test_C1_a_proposal_that_does_not_start_at_its_own_word_is_not_applied(self):
        self.assertEqual(self.applied(fix("text", "accent", "xx", ids(6))), 0)
        self.assertEqual(self.applied(fix("text", "accent", "xx", ids(5, 6))), 0)  # more than one word for a text fix

    def test_C1_an_empty_text_is_only_for_drop_and_a_segment_needs_a_space(self):
        self.assertEqual(self.applied(fix("text", "accent", "  ", ids(5))), 0)
        self.assertEqual(self.applied(fix("segment", "structure", "oneword", ids(5))), 0)
        self.assertEqual(self.applied(fix("drop", "structure", "", ids(5))), 1)

    def test_C3_a_proposal_whose_source_is_a_rule_is_not_applied_by_the_normal_path(self):
        rule = fix("text", "number", "palabraxx", ids(5), source="rule:checks.rule.dni_letter")
        self.assertEqual(self.applied(rule), 0)


class DigitPathTest(DigitTest):
    """The guards of the digit-fix path (A2.5) one by one, and the path through a saved document."""

    def test_C3_a_suggestion_whose_old_digit_is_not_where_it_says_is_stale_and_ignored(self):
        self.assertEqual(self.run_digit(number("p1w50", DNI, 4, "6", "9"))[1], 0)

    def test_C3_a_number_over_several_words_or_a_token_of_a_segment_is_never_fixed(self):
        spread = NumberResult(
            ("p1w50", "p1w51"), DNI + " x", None, "iban", "mismatch", (), DigitSuggestion(4, "5", "9", True, "r")
        )
        self.assertEqual(fixes.apply_digit_fixes([self.page], [spread], "safe")[1], 0)
        self.assertEqual(self.run_digit(number("p1w50#1", DNI, 4, "5", "9"))[1], 0)

    def test_C3_a_stored_rule_fix_is_re_applied_after_a_save_only_if_the_text_is_the_same(self):
        pages, _ = self.run_digit(number("p1w50", DNI, 4, "5", "9"))
        off = fixes.apply_fixes(pages[0], "off")[0]  # the digit fix is forgotten but the proposal stays on the word
        self.assertEqual(off.words[50].text, DNI)
        self.assertEqual(off.words[50].fix.source, "rule:checks.rule.dni_letter")
        bare = replace(off, words=tuple(replace(w, char_alts=None) for w in off.words))  # as after a save and load
        again, count = fixes.apply_digit_fixes([bare], [number("p1w50", DNI, 4, "5", "9")], "safe")
        self.assertEqual((count, again[0].words[50].text), (1, DNI[:4] + "9" + DNI[5:]))
        other, count = fixes.apply_digit_fixes([bare], [number("p1w50", DNI, 4, "5", "7")], "safe")
        self.assertEqual(count, 0)


if __name__ == "__main__":
    unittest.main()
