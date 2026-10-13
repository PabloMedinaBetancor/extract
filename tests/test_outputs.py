import json
import os
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

import jsonschema
import openpyxl

from extract.errors import Cancelled, OutputExists, SourceChanged, UnsupportedFile
from extract.outputs import export as export_module
from extract.outputs import jsondoc
from extract.outputs.export import FORMATS, claim, export
from extract.outputs.text import to_text
from tests.support import SAMPLE, SANS, SPANISH_LINES, sample_document, scanned_pdf, text_page_image

SCHEMA = json.loads((Path(jsondoc.__file__).parent.parent / "schema" / "document.schema.json").read_text())


class Flag:
    def __init__(self, value=False):
        self.value = value

    def is_set(self):
        return self.value


class OutputsTest(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        pages = [text_page_image(SPANISH_LINES, SANS, 24, 1240, 1754) for _ in range(2)]
        self.source = scanned_pdf(self.tmp / "invoice.pdf", pages)
        self.data = sample_document(self.source)
        self.out = self.tmp / "out"


class JsonTest(OutputsTest):
    def test_U1_the_json_validates_against_the_schema(self):
        jsonschema.validate(jsondoc.to_dict(self.data), SCHEMA)

    def test_U1_a_round_trip_gives_an_equal_document(self):
        path = self.tmp / "doc.extract.json"
        jsondoc.save(self.data, path)
        self.assertEqual(jsondoc.load(path), self.data)

    def test_U1_the_schema_refuses_what_the_code_never_writes(self):
        broken = jsondoc.to_dict(self.data)
        broken["pages"][0]["words"][0]["status"] = "fine"
        with self.assertRaises(jsonschema.ValidationError):
            jsonschema.validate(broken, SCHEMA)
        broken = jsondoc.to_dict(self.data)
        broken["numbers"][0]["value"] = 1210.5
        with self.assertRaises(jsonschema.ValidationError):
            jsonschema.validate(broken, SCHEMA)

    def test_U1_exact_decimals_are_strings_and_boxes_come_in_pixels_and_points(self):
        raw = jsondoc.to_dict(self.data)
        self.assertEqual(raw["numbers"][0]["value"], "1000.00")
        word = raw["pages"][0]["words"][0]
        scale = self.data.pages[0].width_pt / self.data.pages[0].width_px
        self.assertAlmostEqual(word["box_pt"][0], word["box_px"][0] * scale)

    def test_U1_fields_carry_the_status_of_their_numbers(self):
        fields = jsondoc.to_dict(self.data)["invoice"]["fields"]
        self.assertEqual(fields["subtotal"]["status"], "checked")
        self.assertEqual(fields["vat"][0]["amount"]["status"], "mismatch")
        self.assertEqual(fields["total"]["status"], "unverified")
        self.assertEqual(fields["number"]["status"], "unverified")

    def test_U1_the_document_says_which_pages_could_not_be_read(self):
        raw = jsondoc.to_dict(self.data)
        self.assertEqual(raw["status"], "partial")
        self.assertEqual(raw["unreadable_pages"], [{"page": 2, "reason": "page.reason.few_words", "error": None}])

    def test_U1_something_else_than_an_extract_document_is_refused(self):
        other = self.tmp / "other.json"
        other.write_text('{"schema_version": 3}')
        with self.assertRaises(UnsupportedFile):
            jsondoc.load(other)
        other.write_text("not json")
        with self.assertRaises(UnsupportedFile):
            jsondoc.load(other)


class EngineDetailTest(OutputsTest):
    """The JSON holds what the user sees; the engine's own detail stays in the cache and never decides equality (R9)."""

    def with_detail(self):
        import dataclasses

        from extract.types import Reading

        page = self.data.pages[0]
        words = list(page.words)
        words[0] = dataclasses.replace(
            words[0],
            char_alts=tuple((("x", 0.9), ("y", 0.1)) for _ in words[0].text),
            masked="masked",
            readings=(Reading("deep", words[0].text, 0.9, None, (), 0, False),),
        )
        page = dataclasses.replace(page, words=tuple(words), rereads={1: ("deep",)}, reread_report={"units_read": 1})
        return dataclasses.replace(self.data, pages=(page, *self.data.pages[1:]))

    def test_R9_the_json_leaves_out_the_engine_detail_and_reads_back_equal(self):
        richer = self.with_detail()
        raw = json.dumps(jsondoc.to_dict(richer))
        for name in ("char_alts", "masked", "readings", "rereads", "reread_report"):
            self.assertNotIn(name, raw)
        path = self.tmp / "detail.extract.json"
        jsondoc.save(richer, path)
        self.assertEqual(jsondoc.load(path), jsondoc.exported(richer))
        export_module._check_json(path, richer)

    def test_R9_documents_that_differ_only_in_engine_detail_are_equal_and_others_are_not(self):
        from extract.api import Document

        richer = self.with_detail()
        self.assertEqual(Document(richer, None), Document(self.data, None))
        import dataclasses

        page = self.data.pages[0]
        changed = dataclasses.replace(page, words=(dataclasses.replace(page.words[0], text="Other"), *page.words[1:]))
        self.assertNotEqual(
            Document(dataclasses.replace(self.data, pages=(changed, *self.data.pages[1:])), None),
            Document(self.data, None),
        )

    def test_C3_a_digit_suggestion_keeps_its_verifying_rule_through_the_json(self):
        import dataclasses

        from extract.types import DigitSuggestion

        number = self.data.numbers[0]
        marked = dataclasses.replace(number, suggestion=DigitSuggestion(1, "0", "8", True, "rule:check_digit"))
        data = dataclasses.replace(self.data, numbers=(marked, *self.data.numbers[1:]))
        raw = jsondoc.to_dict(data)
        jsonschema.validate(raw, SCHEMA)
        self.assertEqual(raw["numbers"][0]["suggestion"]["rule"], "rule:check_digit")
        self.assertEqual(jsondoc.from_dict(raw).numbers[0].suggestion, marked.suggestion)


class ExcelTest(OutputsTest):
    def written(self):
        return export(self.data, self.source, self.out, ["xlsx"])[0]

    def test_U3_numbers_are_numeric_cells_with_their_status_colours(self):
        book = openpyxl.load_workbook(self.written())
        self.assertEqual(book.sheetnames, ["Numbers", "Fields", "Review", "Corrections"])
        numbers = book["Numbers"]
        self.assertEqual([c.value for c in numbers[1]], ["Text", "Value", "Kind", "Status", "Rules", "Page"])
        self.assertIsInstance(numbers["B2"].value, (int, float))
        self.assertEqual(numbers["B2"].value, 1000.0)
        self.assertEqual(numbers["B2"].number_format, "0.00")
        self.assertIsNone(numbers["A2"].fill.fill_type)
        self.assertEqual(numbers["A3"].fill.start_color.rgb[-6:], "FFC7CE")
        self.assertEqual(numbers["A4"].fill.start_color.rgb[-6:], "FFF2CC")
        self.assertEqual(numbers["F2"].value, 1)

    def test_U3_the_review_sheet_lists_what_is_still_doubtful_and_not_what_was_corrected(self):
        review = openpyxl.load_workbook(self.written())["Review"]
        rows = [[c.value for c in row] for row in review.iter_rows(min_row=2)]
        texts = [row[1] for row in rows]
        self.assertIn("2026-0042", texts)
        self.assertIn("€", texts)
        self.assertNotIn(SAMPLE["word"], texts)
        doubtful = next(row for row in rows if row[1] == "2026-0042")
        self.assertEqual(doubtful[2], "2026-0043")
        self.assertIn("read with low confidence", doubtful[3])
        self.assertNotIn("guard.reason", doubtful[3])
        self.assertEqual(next(row for row in rows if row[1] == "€")[3], "not checked")

    def test_U3_the_fields_sheet_has_the_invoice_with_what_was_read(self):
        fields = openpyxl.load_workbook(self.written())["Fields"]
        rows = {row[0].value: [c.value for c in row] for row in fields.iter_rows(min_row=2)}
        self.assertEqual(rows["total"][1], 1210.0)
        self.assertEqual(rows["total"][3], "unverified")
        self.assertEqual(rows["vat.1.amount"][3], "mismatch")
        self.assertEqual(rows["currency"][1], "EUR")

    def test_U3_a_document_without_an_invoice_has_no_fields_sheet(self):
        from dataclasses import replace

        generic = replace(self.data, invoice=None, doc_type="generic")
        path = export(generic, self.source, self.out, ["xlsx"])[0]
        self.assertEqual(openpyxl.load_workbook(path).sheetnames, ["Numbers", "Review", "Corrections"])

    def test_U3_libreoffice_opens_the_workbook(self):
        path = self.written()
        done = subprocess.run(
            ["soffice", "--headless", "--convert-to", "csv", "--outdir", str(self.tmp), str(path)],
            capture_output=True,
            text=True,
            timeout=180,
            env={**os.environ, "HOME": str(self.tmp)},
        )
        csv = self.tmp / (path.stem + ".csv")
        self.assertTrue(csv.exists(), done.stderr)
        self.assertIn("1000", csv.read_text(encoding="utf-8"))


class NeverOverwriteTest(OutputsTest):
    def test_U4_a_second_export_gets_a_numbered_name_and_the_first_stays_as_it_was(self):
        first = export(self.data, self.source, self.out)
        before = {p: p.read_bytes() for p in first}
        second = export(self.data, self.source, self.out)
        third = export(self.data, self.source, self.out)
        self.assertEqual(
            [p.name for p in second],
            ["invoice-2.searchable.pdf", "invoice-2.xlsx", "invoice-2.extract.json", "invoice-2.txt"],
        )
        self.assertEqual([p.name for p in third][0], "invoice-3.searchable.pdf")
        self.assertEqual({p: p.read_bytes() for p in first}, before)

    def test_U4_a_file_that_appears_while_writing_is_never_replaced(self):
        partial = self.tmp / ".x.partial"
        partial.write_text("new")
        final = self.tmp / "x.txt"
        final.write_text("old")
        claimed = claim(partial, self.tmp, "x", ".txt")
        self.assertEqual(claimed.name, "x-2.txt")
        self.assertEqual(final.read_text(), "old")
        self.assertFalse(partial.exists())

    def test_U4_a_failed_check_leaves_nothing_under_the_final_name(self):
        with mock.patch.object(export_module, "_check_json", side_effect=ValueError("bad")):
            with self.assertRaises(ValueError):
                export(self.data, self.source, self.out, ["json"])
        self.assertEqual(sorted(p.name for p in self.out.iterdir()), [])

    def test_U4_a_stop_before_a_file_leaves_no_partial_and_keeps_the_finished_ones_whole(self):
        flag = Flag()
        with self.assertRaises(Cancelled):
            export(
                self.data,
                self.source,
                self.out,
                ["txt", "json"],
                cancel=flag,
                on_written=lambda kind, path: setattr(flag, "value", True),
            )
        self.assertEqual([p.name for p in self.out.iterdir()], ["invoice.txt"])
        self.assertEqual((self.out / "invoice.txt").read_text(encoding="utf-8"), to_text(self.data))

    def test_U4_a_stop_before_everything_writes_nothing(self):
        with self.assertRaises(Cancelled):
            export(self.data, self.source, self.out, cancel=Flag(True))
        self.assertEqual(list(self.out.iterdir()), [])

    def test_A5_a_destination_that_is_a_file_is_refused(self):
        target = self.tmp / "file.txt"
        target.write_text("x")
        with self.assertRaises(OutputExists) as caught:
            export(self.data, self.source, target)
        self.assertEqual(caught.exception.code, "error.output_exists")

    def test_A5_a_source_that_is_not_the_one_read_is_refused(self):
        other = scanned_pdf(self.tmp / "other.pdf", [text_page_image(SPANISH_LINES[:2], SANS, 24, 1240, 1754)])
        with self.assertRaises(SourceChanged) as caught:
            export(self.data, other, self.out, ["pdf"])
        self.assertEqual(caught.exception.code, "error.source_changed")
        self.assertFalse(self.out.exists() and list(self.out.iterdir()))


class ChecksTest(OutputsTest):
    """Each check refuses a bad file: without them a failed write could be given its final name."""

    def test_U4_the_pdf_check_refuses_a_truncated_file_and_one_with_another_page_count(self):
        good = export(self.data, self.source, self.out, ["pdf"])[0]
        truncated = self.tmp / "truncated.pdf"
        truncated.write_bytes(good.read_bytes()[:200])
        with self.assertRaises(Exception):  # noqa: B017 - pikepdf's own error class for a damaged file
            export_module._check_pdf(truncated, self.data)
        short = self.tmp / "short.pdf"
        import pikepdf

        with pikepdf.open(good) as pdf:
            del pdf.pages[1]
            pdf.save(short)
        with self.assertRaises(ValueError):
            export_module._check_pdf(short, self.data)
        export_module._check_pdf(good, self.data)

    def test_U4_the_json_check_refuses_the_json_of_another_document(self):
        from dataclasses import replace

        other = self.tmp / "other.extract.json"
        jsondoc.save(replace(self.data, source="other.pdf"), other)
        with self.assertRaises(ValueError):
            export_module._check_json(other, self.data)
        good = self.tmp / "good.extract.json"
        jsondoc.save(self.data, good)
        export_module._check_json(good, self.data)

    def test_U4_the_workbook_check_refuses_a_file_without_the_numbers_sheet(self):
        book = openpyxl.Workbook()
        book.active.title = "Other"
        wrong = self.tmp / "wrong.xlsx"
        book.save(wrong)
        with self.assertRaises(ValueError):
            export_module._check_xlsx(wrong, self.data)
        with self.assertRaises(Exception):  # noqa: B017 - openpyxl's own error class for a file that is not a workbook
            junk = self.tmp / "junk.xlsx"
            junk.write_bytes(b"not a workbook")
            export_module._check_xlsx(junk, self.data)

    def test_U4_the_text_check_refuses_a_different_text(self):
        wrong = self.tmp / "wrong.txt"
        wrong.write_text("something else", encoding="utf-8")
        with self.assertRaises(ValueError):
            export_module._check_txt(wrong, self.data)

    def test_U4_each_bad_file_is_refused_inside_export_and_leaves_nothing(self):
        for kind, check in (
            ("pdf", "_check_pdf"),
            ("json", "_check_json"),
            ("xlsx", "_check_xlsx"),
            ("txt", "_check_txt"),
        ):
            with self.subTest(kind=kind):
                with mock.patch.object(export_module, check, side_effect=ValueError("bad")):
                    with self.assertRaises(ValueError):
                        export(self.data, self.source, self.out, [kind])
                self.assertEqual(list(self.out.iterdir()), [])


class ReExportTest(OutputsTest):
    def test_U5_exporting_loads_no_engine_and_carries_the_corrected_text(self):
        import pypdfium2 as pdfium

        engine_modules = [m for m in sys.modules if m.startswith("rapidocr")]
        with mock.patch("extract.engine.ocr.load_engine", side_effect=AssertionError("engine loaded")):
            paths = export(self.data, self.source, self.out)
        self.assertEqual([m for m in sys.modules if m.startswith("rapidocr")], engine_modules)
        by_suffix = {p.name.split(".", 1)[1]: p for p in paths}
        self.assertEqual(len(paths), len(FORMATS))
        self.assertIn(SAMPLE["word"], (by_suffix["txt"]).read_text(encoding="utf-8"))
        self.assertNotIn(SAMPLE["word_plain"], by_suffix["txt"].read_text(encoding="utf-8"))
        document = pdfium.PdfDocument(str(by_suffix["searchable.pdf"]))
        page = document[0]
        textpage = page.get_textpage()
        words = textpage.get_text_range().split()
        textpage.close()
        page.close()
        document.close()
        self.assertIn(SAMPLE["word"], words)
        saved = json.loads(by_suffix["extract.json"].read_text(encoding="utf-8"))
        corrected = next(w for w in saved["pages"][0]["words"] if w["review"] == "corrected")
        self.assertEqual(
            (corrected["text"], corrected["original"], corrected["corrected_by"]),
            (SAMPLE["word"], SAMPLE["word_plain"], "user"),
        )

    def test_U5_the_text_has_a_form_feed_between_pages(self):
        pages = to_text(self.data).split("\f")
        self.assertEqual(len(pages), 2)
        self.assertTrue(pages[0].startswith("Factura 2026-0042 Base"))


if __name__ == "__main__":
    unittest.main()


class UnreadableTextLayerTest(OutputsTest):
    def test_U2_pages_given_up_on_that_get_text_in_the_pdf_are_reported(self):
        from dataclasses import replace

        from tests.support import grid_reading

        words = grid_reading(2, SAMPLE["invoice_words"]).words
        pages = list(self.data.pages)
        pages[1] = replace(pages[1], status="unreadable", reason="page.reason.few_words", words=words[:3])
        data = replace(self.data, pages=tuple(pages))
        seen = []
        export(data, self.source, self.out, ["pdf"], on_unreadable_text=seen.append)
        self.assertEqual(seen, [[2]])

    def test_U2_a_page_with_no_words_is_not_reported(self):
        seen = []
        export(self.data, self.source, self.out, ["pdf"], on_unreadable_text=seen.append)
        self.assertEqual(seen, [] if not self.data.pages[1].words else [[2]])
