"""The whole pipeline with the real OCR models: digital, scanned-like and mixed PDFs, and a scanned invoice."""

import tempfile
import unittest
from pathlib import Path

import numpy as np
import openpyxl
import pikepdf
import pypdfium2 as pdfium

import extract
from bench.corpus import invoices
from extract.engine.render import render_page
from tests.fake_guard import FakeGuard
from tests.support import (
    SANS,
    SPANISH_LINES,
    SetLexicon,
    digital_pdf,
    mixed_pdf,
    scanned_pdf,
    text_page_image,
)

LEXICON = SetLexicon({w.lower().strip(".,:") for line in SPANISH_LINES for w in line.split() if len(w) > 2})


class EndToEndTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = Path(tempfile.mkdtemp())
        scan = text_page_image(SPANISH_LINES, SANS, 24, 1240, 1754)
        cls.digital = digital_pdf(cls.tmp / "digital.pdf", [SPANISH_LINES, SPANISH_LINES[:4]])
        cls.scanned = scanned_pdf(cls.tmp / "scanned.pdf", [scan])
        cls.mixed = mixed_pdf(cls.tmp / "mixed.pdf", SPANISH_LINES, scan)
        options = extract.Options(work_dir=cls.tmp / "work", backend="onnxruntime", processes=1, threads=2)
        cls.reader = extract.Reader(options, guard=FakeGuard(), lexicon=LEXICON)
        cls.documents = {p.stem: cls.reader.read(p) for p in (cls.digital, cls.scanned, cls.mixed)}

    @classmethod
    def tearDownClass(cls):
        cls.reader.close()

    def test_A2_each_kind_of_pdf_is_read_with_the_right_page_statuses(self):
        self.assertEqual([p.status for p in self.documents["digital"].pages], ["native", "native"])
        self.assertEqual([p.status for p in self.documents["scanned"].pages], ["ocr"])
        self.assertEqual([p.status for p in self.documents["mixed"].pages], ["native", "ocr"])

    def test_A2_the_scan_reads_as_the_text_that_was_drawn(self):
        expected = " ".join(SPANISH_LINES).split()
        for name in ("scanned", "mixed"):
            document = self.documents[name]
            read = [w.text for w in document.pages[-1].words]
            self.assertGreater(sum(w in read for w in expected) / len(expected), 0.95, name)
            self.assertEqual({w.status for w in document.pages[-1].words}, {"ok"})

    def test_A2_all_four_formats_are_written_for_every_kind_and_reopen(self):
        for name, document in self.documents.items():
            with self.subTest(name=name):
                paths = {p.name.split(".", 1)[1]: p for p in document.export(self.tmp / "out" / name)}
                self.assertEqual(sorted(paths), ["extract.json", "searchable.pdf", "txt", "xlsx"])
                with pikepdf.open(paths["searchable.pdf"]) as pdf:
                    self.assertEqual(len(pdf.pages), len(document.pages))
                self.assertEqual(extract.load(paths["extract.json"], document.source), document)
                self.assertEqual(openpyxl.load_workbook(paths["xlsx"]).sheetnames[0], "Numbers")
                self.assertIn("Factura", paths["txt"].read_text(encoding="utf-8"))

    def test_A2_the_searchable_pdf_of_a_scan_finds_the_words_and_looks_the_same(self):
        document = self.documents["scanned"]
        path = next(p for p in document.export(self.tmp / "search") if p.name.endswith("searchable.pdf"))
        reopened = pdfium.PdfDocument(str(path))
        page = reopened[0]
        textpage = page.get_textpage()
        words = textpage.get_text_range().split()
        textpage.close()
        page.close()
        reopened.close()
        self.assertIn("consultoria", words)
        before, after = pdfium.PdfDocument(str(self.scanned)), pdfium.PdfDocument(str(path))
        try:
            self.assertTrue(np.array_equal(render_page(before[0], 100), render_page(after[0], 100)))
        finally:
            before.close()
            after.close()

    def test_A2_a_scanned_invoice_gives_the_fields_and_numbers_of_the_digital_one(self):
        unit = invoices.plan(5)[0]
        invoice = invoices.make_invoice(unit["id"], unit["layout"], unit["font"], 5)
        digital = self.tmp / "invoice.pdf"
        invoices.render(invoice, unit["layout"], unit["font"], digital)
        document = pdfium.PdfDocument(str(digital))
        image = render_page(document[0], 200)
        document.close()
        scan = scanned_pdf(self.tmp / "invoice-scan.pdf", [image])
        exact = self.reader.read(digital)
        read = self.reader.read(scan)
        self.assertEqual((exact.doc_type, read.doc_type), ("invoice", "invoice"))
        self.assertEqual(read.invoice.total.text, exact.invoice.total.text)
        self.assertEqual(read.invoice.subtotal.text, exact.invoice.subtotal.text)
        self.assertEqual(self.status_of(read, read.invoice.total), self.status_of(exact, exact.invoice.total))
        self.assertEqual(self.status_of(read, read.invoice.total), "checked")

    @staticmethod
    def status_of(document, cell):
        return document.field_status(cell)


if __name__ == "__main__":
    unittest.main()
