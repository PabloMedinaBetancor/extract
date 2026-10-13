"""The Excel workbook: the numbers with their status, the invoice fields, and the words still to review."""

from __future__ import annotations

import json
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font, PatternFill

from ..fixes import corrections
from ..types import DocumentData, NumberResult, Reason, page_of
from .fields import cell_status, named_cells

UNVERIFIED_FILL = PatternFill("solid", start_color="FFF2CC", end_color="FFF2CC")
MISMATCH_FILL = PatternFill("solid", start_color="FFC7CE", end_color="FFC7CE")
MISMATCH_FONT = Font(color="9C0006")
BOLD = Font(bold=True)


def _clean(value: Any) -> Any:
    """A control character (a bell, a null, the 0x1f of a broken font) cannot be stored in a worksheet: openpyxl refuses the
    whole file. It becomes the replacement character, so the cell shows that something was there."""
    return ILLEGAL_CHARACTERS_RE.sub("\ufffd", value) if isinstance(value, str) else value


def _put(sheet, row: list) -> None:
    sheet.append([_clean(value) for value in row])


def load_catalog(language: str) -> dict[str, str]:
    """The texts of the reason and rule keys for a language, from the interface's translation file when it exists."""
    path = Path(__file__).resolve().parent.parent / "ui" / "i18n" / f"{language}.json"
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except FileNotFoundError:
        return {}


def describe(key: str, params: dict[str, Any], catalog: dict[str, str]) -> str:
    template = catalog.get(key)
    if template is None:
        shown = ", ".join(f"{k}={v}" for k, v in params.items())
        return f"{key} ({shown})" if shown else key
    try:
        return template.format(**params)
    except (KeyError, IndexError):
        return template


def _number_cell(value: Decimal | None):
    return None if value is None else float(value)


def _style(cells, status: str) -> None:
    for cell in cells:
        if status == "unverified":
            cell.fill = UNVERIFIED_FILL
        elif status == "mismatch":
            cell.fill = MISMATCH_FILL
            cell.font = MISMATCH_FONT


def _decimals(value: Decimal) -> str:
    places = max(0, -value.as_tuple().exponent) if value.is_finite() else 2
    return "0" if places == 0 else "0." + "0" * places


def _write_numbers(sheet, data: DocumentData, catalog: dict[str, str]) -> None:
    _put(sheet, ["Text", "Value", "Kind", "Status", "Rules", "Page"])
    for number in data.numbers:
        rules = "; ".join(f"{describe(r.key, dict(r.params), catalog)}: {r.result}" for r in number.rules)
        _put(
            sheet,
            [
                number.text,
                _number_cell(number.value),
                number.kind,
                number.status,
                rules,
                page_of(number.word_ids[0]) if number.word_ids else None,
            ],
        )
        row = sheet[sheet.max_row]
        if number.value is not None:
            row[1].number_format = _decimals(number.value)
        _style(row, number.status)


def _as_read(cell, data: DocumentData) -> str:
    originals = {w.id: (w.original if w.original is not None else w.text) for p in data.pages for w in p.words}
    return " ".join(originals.get(i, "") for i in cell.word_ids).strip() or cell.text


def _numbers_by_word(data: DocumentData) -> dict[str, NumberResult]:
    return {i: n for n in data.numbers for i in n.word_ids}


def _write_fields(sheet, data: DocumentData, catalog: dict[str, str]) -> None:
    _put(sheet, ["Field", "Value", "As read", "Status", "Rule", "Page"])
    by_word = _numbers_by_word(data)
    for name, cell in named_cells(data.invoice):
        status = cell_status(cell, data)
        number = next((by_word[i] for i in cell.word_ids if i in by_word), None)
        rule = ""
        if number is not None:
            rule = "; ".join(f"{describe(r.key, dict(r.params), catalog)}: {r.result}" for r in number.rules)
        value = _number_cell(number.value) if number is not None and number.value is not None else cell.text
        _put(sheet, [name, value, _as_read(cell, data), status, rule, cell.page])
        row = sheet[sheet.max_row]
        if isinstance(value, float) and number is not None:
            row[1].number_format = _decimals(number.value)
        _style(row, status)
    if data.invoice.currency is not None:
        _put(sheet, ["currency", data.invoice.currency, data.invoice.currency, "unverified", "", None])
        _style(sheet[sheet.max_row], "unverified")


def _reasons_text(reasons: tuple[Reason, ...], catalog: dict[str, str]) -> str:
    return "; ".join(describe(r.key, dict(r.params), catalog) for r in reasons)


def _write_review(sheet, data: DocumentData, catalog: dict[str, str]) -> None:
    _put(sheet, ["Page", "Text", "Suggestion", "Reasons"])
    for page in data.pages:
        for word in page.words:
            if word.review is not None or word.status not in ("doubtful", "unscored"):
                continue
            reasons = _reasons_text(word.reasons, catalog) if word.status == "doubtful" else "not checked"
            _put(sheet, [page.number, word.text, word.suggestion, reasons])


def _write_corrections(sheet, data: DocumentData) -> None:
    _put(sheet, ["Page", "Operation", "Kind", "Original", "Final", "Source", "Confidence", "Rule", "State"])
    for number, word, state in corrections(data):
        if state == "user":
            _put(sheet, [number, "edit", "user", word.original, word.text, "user", "", "", "applied"])
            continue
        fix = word.fix
        original = word.original if state == "applied" else word.text
        _put(
            sheet,
            [number, fix.op, fix.kind, original, word.text, fix.source, round(fix.p_fix, 4), fix.rule or "", state],
        )


def _no_formulas(book: Workbook) -> None:
    """Nothing in these sheets is a formula. openpyxl stores a string that starts with "=" as one, and the words come
    from a PDF anybody may have made: Excel would run `=HYPERLINK(...)` or `=WEBSERVICE(...)` when the sheet is opened.
    Such a cell is kept as the text it is."""
    for sheet in book.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.data_type == "f":
                    cell.data_type = "s"


def build_workbook(data: DocumentData, language: str = "en") -> Workbook:
    catalog = load_catalog(language)
    book = Workbook()
    numbers = book.active
    numbers.title = "Numbers"
    _write_numbers(numbers, data, catalog)
    if data.invoice is not None:
        _write_fields(book.create_sheet("Fields"), data, catalog)
    _write_review(book.create_sheet("Review"), data, catalog)
    _write_corrections(book.create_sheet("Corrections"), data)
    _no_formulas(book)
    for sheet in book.worksheets:
        for cell in sheet[1]:
            cell.font = BOLD
        sheet.freeze_panes = "A2"
        for column in sheet.columns:
            width = max((len(str(c.value)) for c in column if c.value is not None), default=8)
            sheet.column_dimensions[column[0].column_letter].width = min(max(width + 2, 8), 60)
    return book


def write_excel(data: DocumentData, path: Path, language: str = "en") -> None:
    build_workbook(data, language).save(path)
